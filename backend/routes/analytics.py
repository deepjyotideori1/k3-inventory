from fastapi import APIRouter, HTTPException, Depends, Response
from fastapi.security import HTTPAuthorizationCredentials
from fastapi.responses import StreamingResponse
from database import db
from deps import get_current_user, require_admin, security
from helpers import format_inr, _count_cylinders, _get_date_range_for_period
from typing import List, Optional, Dict, Any
from datetime import datetime, timezone, timedelta
from io import BytesIO
import uuid

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, letter, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
import xlsxwriter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

router = APIRouter()

# ===================== CONNECTION & REFILL ANALYTICS =====================

@router.get("/admin/connection-refill-analytics")
async def get_connection_refill_analytics(
    period: str = 'monthly',
    warehouse_id: Optional[str] = None,
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    credentials: HTTPAuthorizationCredentials = Depends(security)
):
    """Analytics for new connections and refills from sales_entries"""
    user = await get_current_user(credentials)
    if user['role'] != 'admin':
        warehouse_id = user.get('warehouse_id', '')
    
    sd, ed = _get_date_range_for_period(period, start_date, end_date)
    
    query = {'date': {'$gte': sd, '$lte': ed}}
    if warehouse_id and warehouse_id != 'all':
        query['warehouse_id'] = warehouse_id
    
    entries = await db.sales_entries.find(query, {'_id': 0}).to_list(None)
    
    # Categorize
    domestic_new = [e for e in entries if e.get('connection_type') == 'domestic']
    commercial_new = [e for e in entries if e.get('connection_type') == 'commercial']
    domestic_refill = [e for e in entries if e.get('connection_type') == 'domestic_refill']
    commercial_refill = [e for e in entries if e.get('connection_type') == 'commercial_refill']
    
    # Cylinder counts
    domestic_new_cyl = sum(_count_cylinders(e.get('cylinder_nos', '')) for e in domestic_new)
    commercial_new_cyl = sum(_count_cylinders(e.get('cylinder_nos', '')) for e in commercial_new)
    domestic_refill_cyl = sum(int(e.get('no_of_refills', 0) or 0) for e in domestic_refill)
    commercial_refill_cyl = sum(int(e.get('no_of_refills', 0) or 0) for e in commercial_refill)
    
    # For new connections without cylinder_nos, count as 1 per entry
    for e in domestic_new:
        if not e.get('cylinder_nos', '').strip():
            domestic_new_cyl += 1
    for e in commercial_new:
        if not e.get('cylinder_nos', '').strip():
            commercial_new_cyl += 1
    
    # Warehouse breakdown
    warehouses = await db.warehouses.find({}, {'_id': 0}).to_list(None)
    wh_map = {w['id']: w['name'] for w in warehouses}
    
    wh_breakdown = {}
    for e in entries:
        wid = e.get('warehouse_id', '')
        if wid not in wh_breakdown:
            wh_breakdown[wid] = {
                'warehouse_id': wid,
                'warehouse_name': wh_map.get(wid, e.get('warehouse_name', 'Unknown')),
                'domestic_new': 0, 'commercial_new': 0,
                'domestic_new_cyl': 0, 'commercial_new_cyl': 0,
                'domestic_refill': 0, 'commercial_refill': 0,
                'domestic_refill_cyl': 0, 'commercial_refill_cyl': 0
            }
        wb = wh_breakdown[wid]
        ct = e.get('connection_type', '')
        if ct == 'domestic':
            wb['domestic_new'] += 1
            cyl = _count_cylinders(e.get('cylinder_nos', ''))
            wb['domestic_new_cyl'] += cyl if cyl > 0 else 1
        elif ct == 'commercial':
            wb['commercial_new'] += 1
            cyl = _count_cylinders(e.get('cylinder_nos', ''))
            wb['commercial_new_cyl'] += cyl if cyl > 0 else 1
        elif ct == 'domestic_refill':
            wb['domestic_refill'] += 1
            wb['domestic_refill_cyl'] += int(e.get('no_of_refills', 0) or 0)
        elif ct == 'commercial_refill':
            wb['commercial_refill'] += 1
            wb['commercial_refill_cyl'] += int(e.get('no_of_refills', 0) or 0)
    
    # Date-wise breakdown
    from collections import defaultdict
    date_data = defaultdict(lambda: {
        'domestic_new': 0, 'commercial_new': 0,
        'domestic_new_cyl': 0, 'commercial_new_cyl': 0,
        'domestic_refill': 0, 'commercial_refill': 0,
        'domestic_refill_cyl': 0, 'commercial_refill_cyl': 0
    })
    
    for e in entries:
        d = e.get('date', '')
        ct = e.get('connection_type', '')
        dd = date_data[d]
        if ct == 'domestic':
            dd['domestic_new'] += 1
            cyl = _count_cylinders(e.get('cylinder_nos', ''))
            dd['domestic_new_cyl'] += cyl if cyl > 0 else 1
        elif ct == 'commercial':
            dd['commercial_new'] += 1
            cyl = _count_cylinders(e.get('cylinder_nos', ''))
            dd['commercial_new_cyl'] += cyl if cyl > 0 else 1
        elif ct == 'domestic_refill':
            dd['domestic_refill'] += 1
            dd['domestic_refill_cyl'] += int(e.get('no_of_refills', 0) or 0)
        elif ct == 'commercial_refill':
            dd['commercial_refill'] += 1
            dd['commercial_refill_cyl'] += int(e.get('no_of_refills', 0) or 0)
    
    date_breakdown = [{'date': k, **v} for k, v in sorted(date_data.items())]
    
    return {
        'summary': {
            'total_new_connections': len(domestic_new) + len(commercial_new),
            'domestic_new_connections': len(domestic_new),
            'commercial_new_connections': len(commercial_new),
            'domestic_new_cylinders': domestic_new_cyl,
            'commercial_new_cylinders': commercial_new_cyl,
            'total_refills': len(domestic_refill) + len(commercial_refill),
            'domestic_refills': len(domestic_refill),
            'commercial_refills': len(commercial_refill),
            'domestic_refill_cylinders': domestic_refill_cyl,
            'commercial_refill_cylinders': commercial_refill_cyl
        },
        'warehouse_breakdown': list(wh_breakdown.values()),
        'date_breakdown': date_breakdown,
        'period': period,
        'date_range': {'start': sd, 'end': ed}
    }

@router.get("/export/connection-refill-analytics-pdf")
async def export_connection_refill_analytics_pdf(
    period: str = 'monthly',
    warehouse_id: Optional[str] = None,
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    credentials: HTTPAuthorizationCredentials = Depends(security)
):
    """Export connection & refill analytics as PDF"""
    user = await get_current_user(credentials)
    if user['role'] != 'admin':
        warehouse_id = user.get('warehouse_id', '')
    
    sd, ed = _get_date_range_for_period(period, start_date, end_date)
    query = {'date': {'$gte': sd, '$lte': ed}}
    if warehouse_id and warehouse_id != 'all':
        query['warehouse_id'] = warehouse_id
    
    entries = await db.sales_entries.find(query, {'_id': 0}).to_list(None)
    warehouses = await db.warehouses.find({}, {'_id': 0}).to_list(None)
    wh_map = {w['id']: w['name'] for w in warehouses}
    
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), topMargin=30, bottomMargin=30)
    styles = getSampleStyleSheet()
    elements = []
    
    # Title
    title_style = ParagraphStyle('ReportTitle', parent=styles['Title'], fontSize=16, spaceAfter=6)
    elements.append(Paragraph("Connection & Refill Analytics Report", title_style))
    
    wh_label = 'All Warehouses'
    if warehouse_id and warehouse_id != 'all':
        wh_label = wh_map.get(warehouse_id, warehouse_id)
    elements.append(Paragraph(f"Period: {sd} to {ed} | Warehouse: {wh_label}", styles['Normal']))
    elements.append(Spacer(1, 12))
    
    # Summary table
    domestic_new = [e for e in entries if e.get('connection_type') == 'domestic']
    commercial_new = [e for e in entries if e.get('connection_type') == 'commercial']
    domestic_refill = [e for e in entries if e.get('connection_type') == 'domestic_refill']
    commercial_refill = [e for e in entries if e.get('connection_type') == 'commercial_refill']
    
    dn_cyl = sum(_count_cylinders(e.get('cylinder_nos', '')) or 1 for e in domestic_new)
    cn_cyl = sum(_count_cylinders(e.get('cylinder_nos', '')) or 1 for e in commercial_new)
    dr_cyl = sum(int(e.get('no_of_refills', 0) or 0) for e in domestic_refill)
    cr_cyl = sum(int(e.get('no_of_refills', 0) or 0) for e in commercial_refill)
    
    summary_data = [
        ['Category', 'Count', 'Cylinders'],
        ['Domestic New Connections', str(len(domestic_new)), str(dn_cyl)],
        ['Commercial New Connections', str(len(commercial_new)), str(cn_cyl)],
        ['Total New Connections', str(len(domestic_new)+len(commercial_new)), str(dn_cyl+cn_cyl)],
        ['Domestic Refills', str(len(domestic_refill)), str(dr_cyl)],
        ['Commercial Refills', str(len(commercial_refill)), str(cr_cyl)],
        ['Total Refills', str(len(domestic_refill)+len(commercial_refill)), str(dr_cyl+cr_cyl)],
        ['Grand Total', str(len(entries)), str(dn_cyl+cn_cyl+dr_cyl+cr_cyl)]
    ]
    
    st = Table(summary_data, colWidths=[250, 100, 100])
    st.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1e3a5f')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('ALIGN', (1, 0), (-1, -1), 'CENTER'),
        ('BACKGROUND', (0, 3), (-1, 3), colors.HexColor('#e8f5e9')),
        ('BACKGROUND', (0, 6), (-1, 6), colors.HexColor('#e3f2fd')),
        ('BACKGROUND', (0, 7), (-1, 7), colors.HexColor('#fff3e0')),
        ('FONTNAME', (0, 3), (-1, 3), 'Helvetica-Bold'),
        ('FONTNAME', (0, 6), (-1, 6), 'Helvetica-Bold'),
        ('FONTNAME', (0, 7), (-1, 7), 'Helvetica-Bold'),
    ]))
    elements.append(st)
    elements.append(Spacer(1, 16))
    
    # Date-wise breakdown table
    elements.append(Paragraph("Date-wise Breakdown", styles['Heading3']))
    from collections import defaultdict
    date_data = defaultdict(lambda: [0]*8)
    for e in entries:
        d = e.get('date', '')
        ct = e.get('connection_type', '')
        dd = date_data[d]
        if ct == 'domestic':
            dd[0] += 1; dd[1] += _count_cylinders(e.get('cylinder_nos', '')) or 1
        elif ct == 'commercial':
            dd[2] += 1; dd[3] += _count_cylinders(e.get('cylinder_nos', '')) or 1
        elif ct == 'domestic_refill':
            dd[4] += 1; dd[5] += int(e.get('no_of_refills', 0) or 0)
        elif ct == 'commercial_refill':
            dd[6] += 1; dd[7] += int(e.get('no_of_refills', 0) or 0)
    
    dt_table = [['Date', 'Dom. New', 'Cyl', 'Com. New', 'Cyl', 'Dom. Refill', 'Cyl', 'Com. Refill', 'Cyl']]
    grand = [0]*8
    for d in sorted(date_data.keys()):
        row = date_data[d]
        dt_table.append([d] + [str(v) for v in row])
        for i in range(8): grand[i] += row[i]
    dt_table.append(['TOTAL'] + [str(v) for v in grand])
    
    dt = Table(dt_table, colWidths=[70, 55, 40, 55, 40, 60, 40, 60, 40])
    dt.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1e3a5f')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('ALIGN', (1, 0), (-1, -1), 'CENTER'),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#fff3e0')),
        ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
    ]))
    elements.append(dt)
    
    doc.build(elements)
    buffer.seek(0)
    return StreamingResponse(buffer, media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename=connection_refill_analytics_{sd}_to_{ed}.pdf"})

@router.get("/export/connection-refill-analytics-excel")
async def export_connection_refill_analytics_excel(
    period: str = 'monthly',
    warehouse_id: Optional[str] = None,
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    credentials: HTTPAuthorizationCredentials = Depends(security)
):
    """Export connection & refill analytics as Excel"""
    user = await get_current_user(credentials)
    if user['role'] != 'admin':
        warehouse_id = user.get('warehouse_id', '')
    
    sd, ed = _get_date_range_for_period(period, start_date, end_date)
    query = {'date': {'$gte': sd, '$lte': ed}}
    if warehouse_id and warehouse_id != 'all':
        query['warehouse_id'] = warehouse_id
    
    entries = await db.sales_entries.find(query, {'_id': 0}).to_list(None)
    warehouses = await db.warehouses.find({}, {'_id': 0}).to_list(None)
    wh_map = {w['id']: w['name'] for w in warehouses}
    
    output = BytesIO()
    workbook = xlsxwriter.Workbook(output)
    
    header_format = workbook.add_format({'bold': True, 'bg_color': '#1e3a5f', 'font_color': 'white', 'border': 1, 'align': 'center'})
    data_fmt = workbook.add_format({'border': 1, 'align': 'center'})
    bold_fmt = workbook.add_format({'border': 1, 'bold': True, 'align': 'center', 'bg_color': '#fff3e0'})
    green_fmt = workbook.add_format({'border': 1, 'bold': True, 'align': 'center', 'bg_color': '#e8f5e9'})
    blue_fmt = workbook.add_format({'border': 1, 'bold': True, 'align': 'center', 'bg_color': '#e3f2fd'})
    
    # Summary Sheet
    ws1 = workbook.add_worksheet('Summary')
    domestic_new = [e for e in entries if e.get('connection_type') == 'domestic']
    commercial_new = [e for e in entries if e.get('connection_type') == 'commercial']
    domestic_refill = [e for e in entries if e.get('connection_type') == 'domestic_refill']
    commercial_refill = [e for e in entries if e.get('connection_type') == 'commercial_refill']
    
    dn_cyl = sum(_count_cylinders(e.get('cylinder_nos', '')) or 1 for e in domestic_new)
    cn_cyl = sum(_count_cylinders(e.get('cylinder_nos', '')) or 1 for e in commercial_new)
    dr_cyl = sum(int(e.get('no_of_refills', 0) or 0) for e in domestic_refill)
    cr_cyl = sum(int(e.get('no_of_refills', 0) or 0) for e in commercial_refill)
    
    ws1.write(0, 0, f'Connection & Refill Analytics | {sd} to {ed}', workbook.add_format({'bold': True, 'font_size': 14}))
    for col, h in enumerate(['Category', 'Count', 'Cylinders']):
        ws1.write(2, col, h, header_format)
    summary_rows = [
        ('Domestic New Connections', len(domestic_new), dn_cyl),
        ('Commercial New Connections', len(commercial_new), cn_cyl),
        ('Total New Connections', len(domestic_new)+len(commercial_new), dn_cyl+cn_cyl),
        ('Domestic Refills', len(domestic_refill), dr_cyl),
        ('Commercial Refills', len(commercial_refill), cr_cyl),
        ('Total Refills', len(domestic_refill)+len(commercial_refill), dr_cyl+cr_cyl),
        ('Grand Total', len(entries), dn_cyl+cn_cyl+dr_cyl+cr_cyl)
    ]
    for i, (cat, cnt, cyl) in enumerate(summary_rows):
        fmt = green_fmt if i == 2 else (blue_fmt if i == 5 else (bold_fmt if i == 6 else data_fmt))
        ws1.write(3+i, 0, cat, fmt)
        ws1.write(3+i, 1, cnt, fmt)
        ws1.write(3+i, 2, cyl, fmt)
    ws1.set_column(0, 0, 30)
    ws1.set_column(1, 2, 15)
    
    # Date-wise Sheet
    ws2 = workbook.add_worksheet('Date-wise Breakdown')
    headers = ['Date', 'Warehouse', 'Dom. New', 'Dom. New Cyl', 'Com. New', 'Com. New Cyl', 'Dom. Refill', 'Dom. Refill Cyl', 'Com. Refill', 'Com. Refill Cyl']
    for col, h in enumerate(headers):
        ws2.write(0, col, h, header_format)
        ws2.set_column(col, col, 15)
    
    row = 1
    from collections import defaultdict
    date_wh = defaultdict(lambda: [0]*8)
    for e in entries:
        key = (e.get('date', ''), e.get('warehouse_id', ''))
        ct = e.get('connection_type', '')
        dd = date_wh[key]
        if ct == 'domestic':
            dd[0] += 1; dd[1] += _count_cylinders(e.get('cylinder_nos', '')) or 1
        elif ct == 'commercial':
            dd[2] += 1; dd[3] += _count_cylinders(e.get('cylinder_nos', '')) or 1
        elif ct == 'domestic_refill':
            dd[4] += 1; dd[5] += int(e.get('no_of_refills', 0) or 0)
        elif ct == 'commercial_refill':
            dd[6] += 1; dd[7] += int(e.get('no_of_refills', 0) or 0)
    
    grand = [0]*8
    for (d, wid), vals in sorted(date_wh.items()):
        ws2.write(row, 0, d, data_fmt)
        ws2.write(row, 1, wh_map.get(wid, 'Unknown'), data_fmt)
        for i, v in enumerate(vals):
            ws2.write(row, 2+i, v, data_fmt)
            grand[i] += v
        row += 1
    
    ws2.write(row, 0, 'TOTAL', bold_fmt)
    ws2.write(row, 1, '', bold_fmt)
    for i, v in enumerate(grand):
        ws2.write(row, 2+i, v, bold_fmt)
    
    # Warehouse Sheet
    ws3 = workbook.add_worksheet('Warehouse Breakdown')
    wh_headers = ['Warehouse', 'Dom. New', 'Dom. New Cyl', 'Com. New', 'Com. New Cyl', 'Dom. Refill', 'Dom. Refill Cyl', 'Com. Refill', 'Com. Refill Cyl', 'Total']
    for col, h in enumerate(wh_headers):
        ws3.write(0, col, h, header_format)
        ws3.set_column(col, col, 16)
    
    wh_data = defaultdict(lambda: [0]*8)
    for e in entries:
        wid = e.get('warehouse_id', '')
        ct = e.get('connection_type', '')
        dd = wh_data[wid]
        if ct == 'domestic':
            dd[0] += 1; dd[1] += _count_cylinders(e.get('cylinder_nos', '')) or 1
        elif ct == 'commercial':
            dd[2] += 1; dd[3] += _count_cylinders(e.get('cylinder_nos', '')) or 1
        elif ct == 'domestic_refill':
            dd[4] += 1; dd[5] += int(e.get('no_of_refills', 0) or 0)
        elif ct == 'commercial_refill':
            dd[6] += 1; dd[7] += int(e.get('no_of_refills', 0) or 0)
    
    row = 1
    grand = [0]*8
    for wid, vals in sorted(wh_data.items(), key=lambda x: wh_map.get(x[0], '')):
        ws3.write(row, 0, wh_map.get(wid, 'Unknown'), data_fmt)
        total = 0
        for i, v in enumerate(vals):
            ws3.write(row, 1+i, v, data_fmt)
            grand[i] += v
            total += v
        ws3.write(row, 9, total, data_fmt)
        row += 1
    ws3.write(row, 0, 'TOTAL', bold_fmt)
    for i, v in enumerate(grand):
        ws3.write(row, 1+i, v, bold_fmt)
    ws3.write(row, 9, sum(grand), bold_fmt)
    
    workbook.close()
    output.seek(0)
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=connection_refill_analytics_{sd}_to_{ed}.xlsx"})


@router.get("/admin/customer-order-report")
async def get_customer_order_report(
    warehouse_id: Optional[str] = None,
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    search: Optional[str] = None,
    credentials: HTTPAuthorizationCredentials = Depends(security)
):
    """Customer-wise order report grouped by warehouse - all roles with warehouse filtering"""
    user = await get_current_user(credentials)
    
    # Non-admin users are forced to their own warehouse
    if user['role'] != 'admin':
        warehouse_id = user.get('warehouse_id', '')
    
    # Build customer query
    cust_query = {}
    if warehouse_id and warehouse_id != 'all':
        cust_query['warehouse_id'] = warehouse_id
    if search:
        cust_query['$or'] = [
            {'customer_name': {'$regex': search, '$options': 'i'}},
            {'phone': {'$regex': search, '$options': 'i'}},
            {'consumer_no': {'$regex': search, '$options': 'i'}}
        ]
    
    customers = await db.customers.find(cust_query, {'_id': 0}).to_list(5000)
    
    # Get warehouse names
    all_wh = await db.warehouses.find({}, {'_id': 0}).to_list(100)
    wh_map = {w['id']: w['name'] for w in all_wh}
    
    # Build order query
    order_query = {}
    if warehouse_id and warehouse_id != 'all':
        order_query['warehouse_id'] = warehouse_id
    if start_date:
        order_query['order_date'] = order_query.get('order_date', {})
        order_query['order_date']['$gte'] = start_date
    if end_date:
        if 'order_date' not in order_query:
            order_query['order_date'] = {}
        order_query['order_date']['$lte'] = end_date
    
    all_orders = await db.orders.find(order_query, {'_id': 0}).sort('order_date', 1).to_list(10000)
    
    # Also get sales entries for connection/refill history
    sales_query = {}
    if warehouse_id and warehouse_id != 'all':
        sales_query['warehouse_id'] = warehouse_id
    if start_date:
        sales_query['date'] = sales_query.get('date', {})
        sales_query['date']['$gte'] = start_date
    if end_date:
        if 'date' not in sales_query:
            sales_query['date'] = {}
        sales_query['date']['$lte'] = end_date
    
    all_sales = await db.sales_entries.find(sales_query, {'_id': 0}).sort('date', 1).to_list(10000)
    
    # Build customer-id to orders map and name to orders map
    order_by_cid = {}
    order_by_name = {}
    for o in all_orders:
        cid = o.get('customer_id')
        cname = (o.get('customer_name') or '').lower()
        if cid:
            order_by_cid.setdefault(cid, []).append(o)
        elif cname:
            order_by_name.setdefault(cname, []).append(o)
    
    sales_by_cid = {}
    sales_by_name = {}
    for s in all_sales:
        cid = s.get('customer_id')
        cname = (s.get('consumer_name') or '').lower()
        if cid:
            sales_by_cid.setdefault(cid, []).append(s)
        elif cname:
            sales_by_name.setdefault(cname, []).append(s)
    
    total_orders = 0
    total_refills = 0
    customer_groups = []
    
    for c in customers:
        cid = c['id']
        cname = (c.get('customer_name') or '').lower()
        
        # Get orders for this customer
        c_orders = order_by_cid.get(cid, []) + order_by_name.get(cname, [])
        c_sales = sales_by_cid.get(cid, []) + sales_by_name.get(cname, [])
        
        # Combine into unified entries
        entries = []
        seen_ids = set()
        
        for o in c_orders:
            oid = o.get('id', o.get('order_no', ''))
            if oid in seen_ids:
                continue
            seen_ids.add(oid)
            conn_type = o.get('connection_type', '')
            is_refill = 'refill' in conn_type.lower()
            qty = o.get('no_of_cylinders', 1) or 1
            entries.append({
                'source': 'order',
                'id': o.get('order_no', o.get('id', '')),
                'date': o.get('order_date', ''),
                'type': 'Refill' if is_refill else 'New Connection',
                'cylinder_type': conn_type.replace('_refill', '').replace('_', ' ').title(),
                'cylinder_nos': o.get('cylinder_nos', ''),
                'quantity': qty,
                'status': o.get('status', 'pending').title(),
                'payment': o.get('payment_mode', '').replace('_', ' ').title(),
                'memo_no': o.get('memo_no', '')
            })
            total_orders += 1
            if is_refill:
                total_refills += qty
        
        for s in c_sales:
            sid = s.get('id', '')
            if sid in seen_ids:
                continue
            seen_ids.add(sid)
            conn_type = s.get('connection_type', '')
            is_refill = 'refill' in conn_type.lower()
            qty = s.get('no_of_refills', 1) or 1 if is_refill else 1
            entries.append({
                'source': 'sale',
                'id': s.get('memo_no', s.get('id', '')),
                'date': s.get('date', ''),
                'type': 'Refill' if is_refill else 'New Connection',
                'cylinder_type': conn_type.replace('_refill', '').replace('_', ' ').title(),
                'cylinder_nos': s.get('cylinder_nos', ''),
                'quantity': qty,
                'status': 'Completed',
                'payment': s.get('payment_mode', '').replace('_', ' ').title(),
                'memo_no': s.get('memo_no', '')
            })
            total_orders += 1
            if is_refill:
                total_refills += qty
        
        # Sort entries by date ascending
        entries.sort(key=lambda x: x.get('date', ''))
        
        if not entries and not search:
            continue
        
        # Find initial connection date
        connection_date = None
        for e in entries:
            if e['type'] == 'New Connection':
                connection_date = e['date']
                break
        
        customer_groups.append({
            'customer_name': c.get('customer_name', ''),
            'customer_id': c['id'],
            'consumer_no': c.get('consumer_no', ''),
            'phone': c.get('phone', ''),
            'address': c.get('address', ''),
            'connection_type': c.get('connection_type', ''),
            'warehouse_id': c.get('warehouse_id', ''),
            'warehouse_name': wh_map.get(c.get('warehouse_id', ''), 'Unknown'),
            'connection_date': connection_date,
            'total_entries': len(entries),
            'entries': entries
        })
    
    # Sort customer groups by name
    customer_groups.sort(key=lambda x: x['customer_name'].lower())
    
    return {
        'customers': customer_groups,
        'summary': {
            'total_customers': len(customer_groups),
            'total_orders': total_orders,
            'total_refills': total_refills
        }
    }

@router.get("/export/customer-order-report-pdf")
async def export_customer_order_report_pdf(
    warehouse_id: Optional[str] = None,
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    credentials: HTTPAuthorizationCredentials = Depends(security)
):
    """Export customer-wise order report as PDF - all roles"""
    user = await get_current_user(credentials)
    if user['role'] != 'admin':
        warehouse_id = user.get('warehouse_id', '')
    
    # Reuse the report logic
    from starlette.datastructures import QueryParams
    
    # Build same data
    cust_query = {}
    if warehouse_id and warehouse_id != 'all':
        cust_query['warehouse_id'] = warehouse_id
    customers = await db.customers.find(cust_query, {'_id': 0}).to_list(5000)
    
    all_wh = await db.warehouses.find({}, {'_id': 0}).to_list(100)
    wh_map = {w['id']: w['name'] for w in all_wh}
    wh_label = wh_map.get(warehouse_id, 'All Warehouses') if warehouse_id and warehouse_id != 'all' else 'All Warehouses'
    
    order_query = {}
    if warehouse_id and warehouse_id != 'all':
        order_query['warehouse_id'] = warehouse_id
    if start_date:
        order_query['order_date'] = order_query.get('order_date', {})
        order_query['order_date']['$gte'] = start_date
    if end_date:
        if 'order_date' not in order_query:
            order_query['order_date'] = {}
        order_query['order_date']['$lte'] = end_date
    
    all_orders = await db.orders.find(order_query, {'_id': 0}).sort('order_date', 1).to_list(10000)
    
    sales_query = {}
    if warehouse_id and warehouse_id != 'all':
        sales_query['warehouse_id'] = warehouse_id
    if start_date:
        sales_query['date'] = sales_query.get('date', {})
        sales_query['date']['$gte'] = start_date
    if end_date:
        if 'date' not in sales_query:
            sales_query['date'] = {}
        sales_query['date']['$lte'] = end_date
    
    all_sales = await db.sales_entries.find(sales_query, {'_id': 0}).sort('date', 1).to_list(10000)
    
    order_by_cid = {}
    order_by_name = {}
    for o in all_orders:
        cid = o.get('customer_id')
        cname = (o.get('customer_name') or '').lower()
        if cid:
            order_by_cid.setdefault(cid, []).append(o)
        elif cname:
            order_by_name.setdefault(cname, []).append(o)
    
    sales_by_cid = {}
    sales_by_name = {}
    for s in all_sales:
        cid = s.get('customer_id')
        cname = (s.get('consumer_name') or '').lower()
        if cid:
            sales_by_cid.setdefault(cid, []).append(s)
        elif cname:
            sales_by_name.setdefault(cname, []).append(s)
    
    today_display = datetime.now(timezone.utc).strftime('%d-%m-%Y')
    total_orders = 0
    total_refills = 0
    
    output = BytesIO()
    doc = SimpleDocTemplate(output, pagesize=landscape(A4), topMargin=20, bottomMargin=20, leftMargin=20, rightMargin=20)
    elements = []
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle('Title', parent=styles['Title'], fontSize=15, spaceAfter=4, textColor=colors.HexColor('#15803d'))
    subtitle_style = ParagraphStyle('Subtitle', parent=styles['Normal'], fontSize=9, textColor=colors.grey, spaceAfter=6)
    cust_header_style = ParagraphStyle('CustHeader', parent=styles['Heading3'], fontSize=10, textColor=colors.HexColor('#1e3a5f'), spaceBefore=10, spaceAfter=2)
    cust_info_style = ParagraphStyle('CustInfo', parent=styles['Normal'], fontSize=7, textColor=colors.HexColor('#666666'), spaceAfter=4)
    
    elements.append(Paragraph("K3 GAS SERVICE - Customer Order Report", title_style))
    elements.append(Paragraph(f"Warehouse: {wh_label} | Generated: {today_display}", subtitle_style))
    
    for c in sorted(customers, key=lambda x: (x.get('customer_name') or '').lower()):
        cid = c['id']
        cname = (c.get('customer_name') or '').lower()
        c_orders = order_by_cid.get(cid, []) + order_by_name.get(cname, [])
        c_sales = sales_by_cid.get(cid, []) + sales_by_name.get(cname, [])
        
        entries = []
        seen_ids = set()
        for o in c_orders:
            oid = o.get('id', o.get('order_no', ''))
            if oid in seen_ids: continue
            seen_ids.add(oid)
            conn_type = o.get('connection_type', '')
            is_refill = 'refill' in conn_type.lower()
            qty = o.get('no_of_cylinders', 1) or 1
            entries.append({'id': o.get('order_no', ''), 'date': o.get('order_date', ''), 'type': 'Refill' if is_refill else 'Connection', 'cyl_type': conn_type.replace('_refill', '').replace('_', ' ').title(), 'cyl_nos': o.get('cylinder_nos', ''), 'qty': qty, 'status': o.get('status', 'pending').title(), 'payment': o.get('payment_mode', '').title()})
            total_orders += 1
            if is_refill: total_refills += qty
        
        for s in c_sales:
            sid = s.get('id', '')
            if sid in seen_ids: continue
            seen_ids.add(sid)
            conn_type = s.get('connection_type', '')
            is_refill = 'refill' in conn_type.lower()
            qty = s.get('no_of_refills', 1) or 1 if is_refill else 1
            entries.append({'id': s.get('memo_no', sid[:8]), 'date': s.get('date', ''), 'type': 'Refill' if is_refill else 'Connection', 'cyl_type': conn_type.replace('_refill', '').replace('_', ' ').title(), 'cyl_nos': s.get('cylinder_nos', ''), 'qty': qty, 'status': 'Completed', 'payment': s.get('payment_mode', '').title()})
            total_orders += 1
            if is_refill: total_refills += qty
        
        if not entries: continue
        entries.sort(key=lambda x: x.get('date', ''))
        
        wh_name = wh_map.get(c.get('warehouse_id', ''), 'Unknown')
        elements.append(Paragraph(f"{c.get('customer_name', '')} ({wh_name})", cust_header_style))
        elements.append(Paragraph(f"Phone: {c.get('phone', '-')} | Consumer No: {c.get('consumer_no', '-')} | Address: {c.get('address', '-')[:40]}", cust_info_style))
        
        data = [['SL', 'Date', 'Order ID', 'Type', 'Cylinder', 'Cyl Nos', 'Qty', 'Status', 'Payment']]
        for i, e in enumerate(entries, 1):
            try:
                d = datetime.strptime(e['date'], '%Y-%m-%d').strftime('%d-%m-%Y')
            except:
                d = e['date']
            data.append([str(i), d, str(e['id'])[:12], e['type'], e['cyl_type'][:10], e.get('cyl_nos', '')[:10], str(e['qty']), e['status'], e['payment'][:8]])
        
        col_widths = [20, 55, 65, 55, 60, 55, 25, 50, 50]
        table = Table(data, colWidths=col_widths, repeatRows=1)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1e3a5f')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 7),
            ('FONTSIZE', (0, 1), (-1, -1), 7),
            ('TOPPADDING', (0, 0), (-1, -1), 2),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 3),
            ('GRID', (0, 0), (-1, -1), 0.4, colors.grey),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f0f4ff')])
        ]))
        elements.append(table)
        elements.append(Spacer(1, 6))
    
    # Add summary at end
    elements.append(Spacer(1, 12))
    elements.append(Paragraph(f"Summary: {len([c for c in customers])} Customers | {total_orders} Orders | {total_refills} Refills", subtitle_style))
    
    doc.build(elements)
    output.seek(0)
    return StreamingResponse(output, media_type="application/pdf", headers={"Content-Disposition": f"attachment; filename=customer_order_report_{datetime.now(timezone.utc).strftime('%Y-%m-%d')}.pdf"})

@router.get("/export/customer-order-report-excel")
async def export_customer_order_report_excel(
    warehouse_id: Optional[str] = None,
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    credentials: HTTPAuthorizationCredentials = Depends(security)
):
    """Export customer-wise order report as Excel - all roles"""
    user = await get_current_user(credentials)
    if user['role'] != 'admin':
        warehouse_id = user.get('warehouse_id', '')
    
    cust_query = {}
    if warehouse_id and warehouse_id != 'all':
        cust_query['warehouse_id'] = warehouse_id
    customers = await db.customers.find(cust_query, {'_id': 0}).to_list(5000)
    
    all_wh = await db.warehouses.find({}, {'_id': 0}).to_list(100)
    wh_map = {w['id']: w['name'] for w in all_wh}
    wh_label = wh_map.get(warehouse_id, 'All Warehouses') if warehouse_id and warehouse_id != 'all' else 'All Warehouses'
    
    order_query = {}
    if warehouse_id and warehouse_id != 'all':
        order_query['warehouse_id'] = warehouse_id
    if start_date:
        order_query['order_date'] = order_query.get('order_date', {})
        order_query['order_date']['$gte'] = start_date
    if end_date:
        if 'order_date' not in order_query:
            order_query['order_date'] = {}
        order_query['order_date']['$lte'] = end_date
    
    all_orders = await db.orders.find(order_query, {'_id': 0}).sort('order_date', 1).to_list(10000)
    
    sales_query = {}
    if warehouse_id and warehouse_id != 'all':
        sales_query['warehouse_id'] = warehouse_id
    if start_date:
        sales_query['date'] = sales_query.get('date', {})
        sales_query['date']['$gte'] = start_date
    if end_date:
        if 'date' not in sales_query:
            sales_query['date'] = {}
        sales_query['date']['$lte'] = end_date
    
    all_sales = await db.sales_entries.find(sales_query, {'_id': 0}).sort('date', 1).to_list(10000)
    
    order_by_cid = {}
    order_by_name = {}
    for o in all_orders:
        cid = o.get('customer_id')
        cname = (o.get('customer_name') or '').lower()
        if cid:
            order_by_cid.setdefault(cid, []).append(o)
        elif cname:
            order_by_name.setdefault(cname, []).append(o)
    
    sales_by_cid = {}
    sales_by_name = {}
    for s in all_sales:
        cid = s.get('customer_id')
        cname = (s.get('consumer_name') or '').lower()
        if cid:
            sales_by_cid.setdefault(cid, []).append(s)
        elif cname:
            sales_by_name.setdefault(cname, []).append(s)
    
    today_display = datetime.now(timezone.utc).strftime('%d-%m-%Y')
    total_orders = 0
    total_refills = 0
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Customer Order Report"
    
    ws.merge_cells('A1:I1')
    ws['A1'] = "K3 GAS SERVICE - Customer Order Report"
    ws['A1'].font = Font(bold=True, size=14, color="15803d")
    ws.merge_cells('A2:I2')
    ws['A2'] = f"Warehouse: {wh_label} | Generated: {today_display}"
    ws['A2'].font = Font(size=9, color="666666")
    
    current_row = 4
    header_fill = PatternFill(start_color="1e3a5f", end_color="1e3a5f", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=9)
    cust_fill = PatternFill(start_color="e0e7ff", end_color="e0e7ff", fill_type="solid")
    cust_font = Font(bold=True, size=10, color="1e3a5f")
    info_font = Font(size=8, color="666666")
    
    for c in sorted(customers, key=lambda x: (x.get('customer_name') or '').lower()):
        cid = c['id']
        cname = (c.get('customer_name') or '').lower()
        c_orders = order_by_cid.get(cid, []) + order_by_name.get(cname, [])
        c_sales = sales_by_cid.get(cid, []) + sales_by_name.get(cname, [])
        
        entries = []
        seen_ids = set()
        for o in c_orders:
            oid = o.get('id', o.get('order_no', ''))
            if oid in seen_ids: continue
            seen_ids.add(oid)
            conn_type = o.get('connection_type', '')
            is_refill = 'refill' in conn_type.lower()
            qty = o.get('no_of_cylinders', 1) or 1
            entries.append({'id': o.get('order_no', ''), 'date': o.get('order_date', ''), 'type': 'Refill' if is_refill else 'Connection', 'cyl_type': conn_type.replace('_refill', '').replace('_', ' ').title(), 'cyl_nos': o.get('cylinder_nos', ''), 'qty': qty, 'status': o.get('status', 'pending').title(), 'payment': o.get('payment_mode', '').title()})
            total_orders += 1
            if is_refill: total_refills += qty
        
        for s in c_sales:
            sid = s.get('id', '')
            if sid in seen_ids: continue
            seen_ids.add(sid)
            conn_type = s.get('connection_type', '')
            is_refill = 'refill' in conn_type.lower()
            qty = s.get('no_of_refills', 1) or 1 if is_refill else 1
            entries.append({'id': s.get('memo_no', sid[:8]), 'date': s.get('date', ''), 'type': 'Refill' if is_refill else 'Connection', 'cyl_type': conn_type.replace('_refill', '').replace('_', ' ').title(), 'cyl_nos': s.get('cylinder_nos', ''), 'qty': qty, 'status': 'Completed', 'payment': s.get('payment_mode', '').title()})
            total_orders += 1
            if is_refill: total_refills += qty
        
        if not entries: continue
        entries.sort(key=lambda x: x.get('date', ''))
        
        wh_name = wh_map.get(c.get('warehouse_id', ''), 'Unknown')
        
        # Customer header row
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=9)
        cell = ws.cell(row=current_row, column=1, value=f"{c.get('customer_name', '')} ({wh_name})")
        cell.fill = cust_fill
        cell.font = cust_font
        current_row += 1
        
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=9)
        cell = ws.cell(row=current_row, column=1, value=f"Phone: {c.get('phone', '-')} | Consumer No: {c.get('consumer_no', '-')} | Address: {c.get('address', '-')}")
        cell.font = info_font
        current_row += 1
        
        headers = ['SL', 'Date', 'Order ID', 'Type', 'Cylinder Type', 'Cyl Nos', 'Qty', 'Status', 'Payment']
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=current_row, column=col, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center')
        current_row += 1
        
        for i, e in enumerate(entries, 1):
            try:
                d = datetime.strptime(e['date'], '%Y-%m-%d').strftime('%d-%m-%Y')
            except:
                d = e['date']
            ws.append([i, d, str(e['id']), e['type'], e['cyl_type'], e.get('cyl_nos', ''), e['qty'], e['status'], e['payment']])
            current_row += 1
        
        current_row += 1
    
    # Summary
    ws.cell(row=current_row + 1, column=1, value=f"Total Customers: {len(customers)} | Total Orders: {total_orders} | Total Refills: {total_refills}").font = Font(bold=True, size=10)
    
    col_widths = [6, 12, 16, 12, 14, 12, 6, 12, 12]
    for idx, w in enumerate(col_widths, 1):
        ws.column_dimensions[chr(64 + idx)].width = w
    
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f"attachment; filename=customer_order_report_{datetime.now(timezone.utc).strftime('%Y-%m-%d')}.xlsx"})

@router.get("/orders/summary/stats")
async def get_order_summary(
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    credentials: HTTPAuthorizationCredentials = Depends(security)
):
    """Get order summary statistics"""
    user = await get_current_user(credentials)
    
    query = {}
    if user['role'] != 'admin':
        query['warehouse_id'] = user.get('warehouse_id')
    
    if start_date:
        query['order_date'] = query.get('order_date', {})
        query['order_date']['$gte'] = start_date
    if end_date:
        if 'order_date' not in query:
            query['order_date'] = {}
        query['order_date']['$lte'] = end_date
    
    # Get counts
    total_orders = await db.orders.count_documents(query)
    
    domestic_query = {**query, 'connection_type': 'domestic'}
    commercial_query = {**query, 'connection_type': 'commercial'}
    cash_query = {**query, 'payment_mode': 'cash'}
    online_query = {**query, 'payment_mode': 'online'}
    credit_query = {**query, 'payment_mode': 'credit_pending'}
    pending_query = {**query, '$or': [{'status': 'pending'}, {'status': {'$exists': False}}]}
    delivered_query = {**query, 'status': 'delivered'}
    
    total_domestic = await db.orders.count_documents(domestic_query)
    total_commercial = await db.orders.count_documents(commercial_query)
    total_cash = await db.orders.count_documents(cash_query)
    total_online = await db.orders.count_documents(online_query)
    total_credit = await db.orders.count_documents(credit_query)
    total_pending = await db.orders.count_documents(pending_query)
    total_delivered = await db.orders.count_documents(delivered_query)
    cancelled_query = {**query, 'status': 'cancelled'}
    total_cancelled = await db.orders.count_documents(cancelled_query)
    
    return {
        'total_orders': total_orders,
        'total_domestic': total_domestic,
        'total_commercial': total_commercial,
        'total_cash': total_cash,
        'total_online': total_online,
        'total_credit_pending': total_credit,
        'total_pending': total_pending,
        'total_delivered': total_delivered,
        'total_cancelled': total_cancelled
    }

@router.get("/orders/pdf/{order_id}")
async def download_order_pdf(
    order_id: str,
    credentials: HTTPAuthorizationCredentials = Depends(security)
):
    """Download a single order as PDF"""
    user = await get_current_user(credentials)
    
    order = await db.orders.find_one({'id': order_id})
    if not order:
        raise HTTPException(status_code=404, detail="Order not found")
    
    # Check access
    if user['role'] != 'admin' and order.get('warehouse_id') != user.get('warehouse_id'):
        raise HTTPException(status_code=403, detail="Access denied")
    
    warehouse = await db.warehouses.find_one({'id': order.get('warehouse_id')})
    warehouse_name = warehouse['name'] if warehouse else 'Unknown'
    
    # Create PDF
    output = BytesIO()
    doc = SimpleDocTemplate(output, pagesize=A4, topMargin=30, bottomMargin=30)
    elements = []
    styles = getSampleStyleSheet()
    
    # Title style
    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        fontSize=18,
        textColor=colors.HexColor('#2d5016'),
        spaceAfter=10,
        alignment=1
    )
    
    subtitle_style = ParagraphStyle(
        'Subtitle',
        parent=styles['Normal'],
        fontSize=12,
        textColor=colors.grey,
        alignment=1,
        spaceAfter=20
    )
    
    # Header
    elements.append(Paragraph("K3 GAS SERVICE", title_style))
    elements.append(Paragraph("Khayal Hamesha", subtitle_style))
    elements.append(Spacer(1, 10))
    
    # Order details box
    order_title = ParagraphStyle('OrderTitle', parent=styles['Heading2'], fontSize=14, textColor=colors.HexColor('#2d5016'))
    elements.append(Paragraph(f"ORDER: {order['order_no']}", order_title))
    elements.append(Spacer(1, 10))
    
    # Order info table
    order_data = [
        ['Order Date:', order['order_date']],
        ['Order No:', order['order_no']],
        ['Warehouse:', warehouse_name],
        ['Customer Name:', order['customer_name']],
        ['Mobile Number:', order.get('mobile_number', '-')],
        ['Address/Landmark:', order.get('address_landmark', '-')],
        ['Connection Type:', order['connection_type'].capitalize()],
        ['Payment Mode:', order['payment_mode'].replace('_', ' ').title()],
        ['Remarks:', order.get('remarks', '-')],
    ]
    
    table = Table(order_data, colWidths=[150, 300])
    table.setStyle(TableStyle([
        ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 11),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
        ('TOPPADDING', (0, 0), (-1, -1), 8),
        ('ALIGN', (0, 0), (0, -1), 'RIGHT'),
        ('ALIGN', (1, 0), (1, -1), 'LEFT'),
        ('TEXTCOLOR', (0, 0), (0, -1), colors.HexColor('#666666')),
        ('LINEBELOW', (0, 0), (-1, -2), 0.5, colors.HexColor('#eeeeee')),
    ]))
    elements.append(table)
    
    elements.append(Spacer(1, 30))
    
    # Footer
    footer_style = ParagraphStyle('Footer', parent=styles['Normal'], fontSize=9, textColor=colors.grey, alignment=1)
    elements.append(Paragraph(f"Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M')}", footer_style))
    
    doc.build(elements)
    output.seek(0)
    
    return Response(
        content=output.getvalue(),
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename=Order_{order['order_no']}_{datetime.now().strftime('%d%m%y')}.pdf"}
    )

@router.get("/export/orders-pdf")
async def export_orders_pdf(
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    payment_mode: Optional[str] = None,
    connection_type: Optional[str] = None,
    credentials: HTTPAuthorizationCredentials = Depends(security)
):
    """Export orders report to PDF"""
    user = await get_current_user(credentials)
    
    query = {}
    if user['role'] != 'admin':
        query['warehouse_id'] = user.get('warehouse_id')
    
    if start_date:
        query['order_date'] = query.get('order_date', {})
        query['order_date']['$gte'] = start_date
    if end_date:
        if 'order_date' not in query:
            query['order_date'] = {}
        query['order_date']['$lte'] = end_date
    
    if payment_mode and payment_mode != 'all':
        query['payment_mode'] = payment_mode
    if connection_type and connection_type != 'all':
        query['connection_type'] = connection_type
    
    orders = await db.orders.find(query).sort('order_date', -1).to_list(1000)
    
    # Get warehouse name
    warehouse_name = "All Warehouses"
    if user['role'] != 'admin':
        warehouse = await db.warehouses.find_one({'id': user.get('warehouse_id')})
        warehouse_name = warehouse['name'] if warehouse else 'Unknown'
    
    # Create PDF - A4 landscape fit-to-page
    output = BytesIO()
    doc = SimpleDocTemplate(output, pagesize=landscape(A4), topMargin=15, bottomMargin=15, leftMargin=15, rightMargin=15)
    elements = []
    styles = getSampleStyleSheet()
    
    # Header 14pt bold
    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        fontSize=14,
        textColor=colors.HexColor('#2d5016'),
        spaceAfter=5,
        alignment=1
    )
    
    date_range = ""
    if start_date and end_date:
        date_range = f" ({start_date} to {end_date})"
    elif start_date:
        date_range = f" (from {start_date})"
    elif end_date:
        date_range = f" (until {end_date})"
    
    elements.append(Paragraph(f"K3 GAS SERVICE - Orders Report{date_range}", title_style))
    elements.append(Paragraph(f"Warehouse: {warehouse_name}", ParagraphStyle('Sub', fontSize=10, alignment=1)))
    elements.append(Spacer(1, 5))
    
    # Table data with clear headers
    table_data = [['Date', 'Order No', 'Customer', 'Mobile', 'Address', 'Type', 'Payment', 'Status', 'Remarks']]
    
    for o in orders:
        status = o.get('status', 'pending').title()
        if status == 'Cancelled' and o.get('cancellation_reason'):
            status = f"Cancelled: {o.get('cancellation_reason', '')[:15]}"
        table_data.append([
            o.get('order_date', ''),
            o.get('order_no', ''),
            o.get('customer_name', '')[:18],
            o.get('mobile_number', ''),
            o.get('address_landmark', '')[:20],
            o.get('connection_type', '').replace('_', ' ').title()[:10],
            o.get('payment_mode', '').title()[:6],
            status[:18],
            o.get('remarks', '')[:12]
        ])
    
    # Fit to A4 landscape
    col_widths = [55, 45, 100, 70, 110, 70, 50, 75, 70]
    table = Table(table_data, colWidths=col_widths, repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2d5016')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 8),
        ('FONTSIZE', (0, 1), (-1, -1), 7),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 4),
        ('TOPPADDING', (0, 0), (-1, -1), 2),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f0f0f0')])
    ]))
    
    elements.append(table)
    elements.append(Spacer(1, 10))
    elements.append(Paragraph(f"Total Orders: {len(orders)}", ParagraphStyle('Total', fontSize=10)))
    
    doc.build(elements)
    output.seek(0)
    
    return Response(
        content=output.getvalue(),
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename=Orders_Report_{datetime.now().strftime('%d%m%y')}.pdf"}
    )

@router.get("/export/orders-excel")
async def export_orders_excel(
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    payment_mode: Optional[str] = None,
    connection_type: Optional[str] = None,
    credentials: HTTPAuthorizationCredentials = Depends(security)
):
    """Export orders report to Excel"""
    user = await get_current_user(credentials)
    
    query = {}
    if user['role'] != 'admin':
        query['warehouse_id'] = user.get('warehouse_id')
    
    if start_date:
        query['order_date'] = query.get('order_date', {})
        query['order_date']['$gte'] = start_date
    if end_date:
        if 'order_date' not in query:
            query['order_date'] = {}
        query['order_date']['$lte'] = end_date
    
    if payment_mode and payment_mode != 'all':
        query['payment_mode'] = payment_mode
    if connection_type and connection_type != 'all':
        query['connection_type'] = connection_type
    
    orders = await db.orders.find(query).sort('order_date', -1).to_list(1000)
    
    # Get warehouse names
    warehouse_ids = list(set(o.get('warehouse_id') for o in orders if o.get('warehouse_id')))
    warehouses = await db.warehouses.find({'id': {'$in': warehouse_ids}}).to_list(100)
    warehouse_map = {w['id']: w['name'] for w in warehouses}
    
    # Create Excel
    output = BytesIO()
    workbook = xlsxwriter.Workbook(output)
    worksheet = workbook.add_worksheet('Orders')
    
    # Formats
    header_format = workbook.add_format({
        'bold': True,
        'bg_color': '#2d5016',
        'font_color': 'white',
        'border': 1,
        'align': 'center'
    })
    
    data_format = workbook.add_format({'border': 1, 'align': 'left'})
    cash_format = workbook.add_format({'border': 1, 'align': 'center', 'bg_color': '#d4edda'})
    online_format = workbook.add_format({'border': 1, 'align': 'center', 'bg_color': '#cce5ff'})
    credit_format = workbook.add_format({'border': 1, 'align': 'center', 'bg_color': '#fff3cd'})
    cancelled_format = workbook.add_format({'border': 1, 'align': 'center', 'bg_color': '#f8d7da', 'font_color': '#721c24'})
    delivered_format = workbook.add_format({'border': 1, 'align': 'center', 'bg_color': '#d4edda', 'font_color': '#155724'})
    pending_format = workbook.add_format({'border': 1, 'align': 'center', 'bg_color': '#fff3cd', 'font_color': '#856404'})
    
    # Headers
    headers = ['Order Date', 'Order No', 'Customer Name', 'Mobile', 'Address/Landmark', 'Connection Type', 'Payment Mode', 'Status', 'Remarks', 'Warehouse']
    
    for col, header in enumerate(headers):
        worksheet.write(0, col, header, header_format)
        worksheet.set_column(col, col, 15 if col < 3 else 20)
    
    # Data
    for row, o in enumerate(orders, start=1):
        worksheet.write(row, 0, o.get('order_date', ''), data_format)
        worksheet.write(row, 1, o.get('order_no', ''), data_format)
        worksheet.write(row, 2, o.get('customer_name', ''), data_format)
        worksheet.write(row, 3, o.get('mobile_number', ''), data_format)
        worksheet.write(row, 4, o.get('address_landmark', ''), data_format)
        worksheet.write(row, 5, o.get('connection_type', '').capitalize(), data_format)
        
        pm = o.get('payment_mode', '')
        pm_format = cash_format if pm == 'cash' else (online_format if pm == 'online' else credit_format)
        worksheet.write(row, 6, pm.replace('_', ' ').title(), pm_format)
        
        status = o.get('status', 'pending')
        status_label = status.title()
        if status == 'cancelled' and o.get('cancellation_reason'):
            status_label = f"Cancelled: {o.get('cancellation_reason', '')}"
        s_fmt = cancelled_format if status == 'cancelled' else (delivered_format if status == 'delivered' else pending_format)
        worksheet.write(row, 7, status_label, s_fmt)
        
        worksheet.write(row, 8, o.get('remarks', ''), data_format)
        worksheet.write(row, 9, warehouse_map.get(o.get('warehouse_id', ''), 'Unknown'), data_format)
    
    workbook.close()
    output.seek(0)
    
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=Orders_Report_{datetime.now().strftime('%d%m%y')}.xlsx"}
    )

