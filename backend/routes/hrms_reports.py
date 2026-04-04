from fastapi import APIRouter, HTTPException, Depends, Response
from database import db
from deps import get_current_user
from helpers import format_inr
from datetime import datetime, timezone
from typing import Optional
from io import BytesIO
import calendar

router = APIRouter()


def get_company_info_sync(settings):
    return {
        'name': settings.get('company_name', 'K3 GAS SERVICE') if settings else 'K3 GAS SERVICE',
        'tagline': settings.get('tagline', '') if settings else '',
        'address': settings.get('address', '') if settings else '',
        'email': settings.get('email', '') if settings else '',
        'helpline': settings.get('helpline', '') if settings else '',
    }


def build_pdf_header(elements, company, title, styles, colors, Paragraph, Spacer, Table, TableStyle, TA_CENTER, mm):
    """Reusable A4 report header with company info"""
    title_style = styles['Title']
    title_style.fontName = 'Helvetica-Bold'
    title_style.fontSize = 14
    title_style.alignment = TA_CENTER

    elements.append(Paragraph(company['name'], title_style))
    if company['address']:
        elements.append(Paragraph(company['address'], styles['Normal']))
    elements.append(Spacer(1, 3 * mm))
    elements.append(Paragraph(title, styles['Heading2']))
    elements.append(Spacer(1, 4 * mm))


# ============ EMPLOYEE LIST REPORT (PDF) ============

@router.get("/hrms/reports/employees/pdf")
async def export_employees_pdf(
    department_id: Optional[str] = None,
    status: Optional[str] = None,
    user: dict = Depends(get_current_user)
):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.lib.enums import TA_CENTER, TA_LEFT

    settings = await db.hrms_settings.find_one({'key': 'company_info'}, {'_id': 0})
    company = get_company_info_sync(settings)

    query = {}
    if department_id and department_id != 'all':
        query['department_id'] = department_id
    if status == 'active':
        query['is_active'] = True
    elif status == 'inactive':
        query['is_active'] = False

    employees = []
    async for emp in db.hrms_employees.find(query, {'_id': 0}).sort('name', 1):
        dept = await db.hrms_departments.find_one({'id': emp.get('department_id')}, {'_id': 0, 'name': 1})
        emp['department_name'] = dept['name'] if dept else 'N/A'
        employees.append(emp)

    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=12 * mm, rightMargin=12 * mm, topMargin=15 * mm, bottomMargin=15 * mm)
    styles = getSampleStyleSheet()
    elements = []

    build_pdf_header(elements, company, "EMPLOYEE DIRECTORY", styles, colors, Paragraph, Spacer, Table, TableStyle, TA_CENTER, mm)

    elements.append(Paragraph(f"Total Employees: {len(employees)} | Generated: {datetime.now().strftime('%d-%m-%Y %H:%M')}", ParagraphStyle('Info', fontName='Helvetica', fontSize=8, alignment=TA_CENTER, textColor=colors.grey)))
    elements.append(Spacer(1, 4 * mm))

    data = [['S.No', 'Emp ID', 'Name', 'Department', 'Designation', 'DOJ', 'Phone', 'Status']]
    for i, emp in enumerate(employees):
        data.append([
            str(i + 1),
            emp.get('employee_id', ''),
            emp['name'],
            emp.get('department_name', ''),
            emp.get('designation', ''),
            emp.get('date_of_joining', ''),
            emp.get('phone', ''),
            'Active' if emp.get('is_active') else 'Inactive',
        ])

    table = Table(data, colWidths=[25, 45, 80, 65, 70, 55, 65, 40])
    table.setStyle(TableStyle([
        ('FONTNAME', (0, 0), (-1, -1), 'Helvetica'),
        ('FONTSIZE', (0, 0), (-1, -1), 7),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 8),
        ('BACKGROUND', (0, 0), (-1, 0), colors.Color(0.15, 0.3, 0.6)),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ALIGN', (0, 0), (0, -1), 'CENTER'),
        ('ALIGN', (-1, 0), (-1, -1), 'CENTER'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.lightgrey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.Color(0.97, 0.97, 0.97)]),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
    ]))
    elements.append(table)

    doc.build(elements)
    buf.seek(0)
    return Response(
        content=buf.read(),
        media_type="application/pdf",
        headers={"Content-Disposition": 'attachment; filename="Employee_Directory.pdf"'}
    )


# ============ EMPLOYEE LIST REPORT (EXCEL) ============

@router.get("/hrms/reports/employees/excel")
async def export_employees_excel(
    department_id: Optional[str] = None,
    status: Optional[str] = None,
    user: dict = Depends(get_current_user)
):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

    settings = await db.hrms_settings.find_one({'key': 'company_info'}, {'_id': 0})
    company = get_company_info_sync(settings)

    query = {}
    if department_id and department_id != 'all':
        query['department_id'] = department_id
    if status == 'active':
        query['is_active'] = True
    elif status == 'inactive':
        query['is_active'] = False

    employees = []
    async for emp in db.hrms_employees.find(query, {'_id': 0}).sort('name', 1):
        dept = await db.hrms_departments.find_one({'id': emp.get('department_id')}, {'_id': 0, 'name': 1})
        emp['department_name'] = dept['name'] if dept else 'N/A'
        employees.append(emp)

    wb = Workbook()
    ws = wb.active
    ws.title = "Employees"

    header_font = Font(name='Arial', size=14, bold=True)
    sub_font = Font(name='Arial', size=9, color='666666')
    col_font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
    cell_font = Font(name='Arial', size=9)
    header_fill = PatternFill(start_color='264785', end_color='264785', fill_type='solid')
    thin_border = Border(
        left=Side(style='thin', color='CCCCCC'),
        right=Side(style='thin', color='CCCCCC'),
        top=Side(style='thin', color='CCCCCC'),
        bottom=Side(style='thin', color='CCCCCC'),
    )

    ws.merge_cells('A1:J1')
    ws['A1'] = company['name']
    ws['A1'].font = header_font
    ws['A1'].alignment = Alignment(horizontal='center')

    ws.merge_cells('A2:J2')
    ws['A2'] = f"Employee Directory | Generated: {datetime.now().strftime('%d-%m-%Y')}"
    ws['A2'].font = sub_font
    ws['A2'].alignment = Alignment(horizontal='center')

    headers = ['S.No', 'Emp ID', 'Name', 'Department', 'Designation', 'DOJ', 'Phone', 'Email', 'PAN', 'Status']
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col, value=h)
        cell.font = col_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
        cell.border = thin_border

    for i, emp in enumerate(employees):
        row = i + 5
        values = [
            i + 1, emp.get('employee_id', ''), emp['name'], emp.get('department_name', ''),
            emp.get('designation', ''), emp.get('date_of_joining', ''), emp.get('phone', ''),
            emp.get('email', ''), emp.get('pan_number', ''),
            'Active' if emp.get('is_active') else 'Inactive'
        ]
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=val)
            cell.font = cell_font
            cell.border = thin_border

    for col in [('A', 6), ('B', 10), ('C', 20), ('D', 16), ('E', 18), ('F', 12), ('G', 14), ('H', 24), ('I', 12), ('J', 10)]:
        ws.column_dimensions[col[0]].width = col[1]

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return Response(
        content=buf.read(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="Employee_Directory.xlsx"'}
    )


# ============ ATTENDANCE REPORT (PDF) ============

@router.get("/hrms/reports/attendance/pdf")
async def export_attendance_pdf(month: str, department_id: Optional[str] = None, user: dict = Depends(get_current_user)):
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.lib.enums import TA_CENTER

    settings = await db.hrms_settings.find_one({'key': 'company_info'}, {'_id': 0})
    company = get_company_info_sync(settings)
    parts = month.split('-')
    year_num, month_num = int(parts[0]), int(parts[1])
    month_name = calendar.month_name[month_num]

    emp_query = {'is_active': True}
    if department_id and department_id != 'all':
        emp_query['department_id'] = department_id

    summaries = []
    async for emp in db.hrms_employees.find(emp_query, {'_id': 0}).sort('name', 1):
        dept = await db.hrms_departments.find_one({'id': emp.get('department_id')}, {'_id': 0, 'name': 1})
        present = await db.hrms_attendance.count_documents({'employee_id': emp['id'], 'date': {'$regex': f'^{month}'}, 'status': 'present'})
        absent = await db.hrms_attendance.count_documents({'employee_id': emp['id'], 'date': {'$regex': f'^{month}'}, 'status': 'absent'})
        half_day = await db.hrms_attendance.count_documents({'employee_id': emp['id'], 'date': {'$regex': f'^{month}'}, 'status': 'half_day'})
        late = await db.hrms_attendance.count_documents({'employee_id': emp['id'], 'date': {'$regex': f'^{month}'}, 'status': 'late'})
        leave = await db.hrms_attendance.count_documents({'employee_id': emp['id'], 'date': {'$regex': f'^{month}'}, 'status': 'leave'})
        effective = present + late + (half_day * 0.5)
        summaries.append({
            'code': emp.get('employee_id', ''), 'name': emp['name'],
            'dept': dept['name'] if dept else 'N/A', 'present': present,
            'absent': absent, 'half_day': half_day, 'late': late,
            'leave': leave, 'effective': effective,
        })

    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm, topMargin=15 * mm, bottomMargin=15 * mm)
    styles = getSampleStyleSheet()
    elements = []

    build_pdf_header(elements, company, f"ATTENDANCE REPORT - {month_name} {year_num}", styles, colors, Paragraph, Spacer, Table, TableStyle, TA_CENTER, mm)

    data = [['S.No', 'Emp ID', 'Name', 'Department', 'Present', 'Absent', 'Half Day', 'Late', 'Leave', 'Effective Days']]
    for i, s in enumerate(summaries):
        data.append([str(i + 1), s['code'], s['name'], s['dept'], str(s['present']), str(s['absent']), str(s['half_day']), str(s['late']), str(s['leave']), str(s['effective'])])

    totals = ['', '', f'Total ({len(summaries)})', '', str(sum(s['present'] for s in summaries)), str(sum(s['absent'] for s in summaries)),
              str(sum(s['half_day'] for s in summaries)), str(sum(s['late'] for s in summaries)), str(sum(s['leave'] for s in summaries)),
              str(sum(s['effective'] for s in summaries))]
    data.append(totals)

    table = Table(data, colWidths=[25, 45, 90, 70, 45, 45, 45, 40, 40, 55])
    table.setStyle(TableStyle([
        ('FONTNAME', (0, 0), (-1, -1), 'Helvetica'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('BACKGROUND', (0, 0), (-1, 0), colors.Color(0.15, 0.3, 0.6)),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ALIGN', (4, 0), (-1, -1), 'CENTER'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.lightgrey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -2), [colors.white, colors.Color(0.97, 0.97, 0.97)]),
        ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
        ('BACKGROUND', (0, -1), (-1, -1), colors.Color(0.9, 0.92, 0.95)),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
    ]))
    elements.append(table)
    doc.build(elements)
    buf.seek(0)
    return Response(content=buf.read(), media_type="application/pdf",
                    headers={"Content-Disposition": f'attachment; filename="Attendance_{month_name}_{year_num}.pdf"'})


# ============ PAYROLL REPORT (PDF) ============

@router.get("/hrms/reports/payroll/{payroll_id}/pdf")
async def export_payroll_report_pdf(payroll_id: str, user: dict = Depends(get_current_user)):
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.lib.enums import TA_CENTER

    payroll = await db.hrms_payroll.find_one({'id': payroll_id}, {'_id': 0})
    if not payroll:
        raise HTTPException(status_code=404, detail="Payroll not found")

    settings = await db.hrms_settings.find_one({'key': 'company_info'}, {'_id': 0})
    company = get_company_info_sync(settings)
    month_name = calendar.month_name[payroll['month']]

    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm, topMargin=12 * mm, bottomMargin=12 * mm)
    styles = getSampleStyleSheet()
    elements = []

    build_pdf_header(elements, company, f"PAYROLL REGISTER - {month_name} {payroll['year']}", styles, colors, Paragraph, Spacer, Table, TableStyle, TA_CENTER, mm)

    data = [['S.No', 'Emp ID', 'Name', 'Dept', 'Days', 'Basic', 'HRA', 'DA', 'Other', 'Gross', 'PF', 'ESI', 'PT', 'TDS', 'Deductions', 'Net Pay']]
    for i, emp in enumerate(payroll.get('employees', [])):
        data.append([
            str(i + 1), emp.get('employee_code', ''), emp['name'][:18], emp.get('department', '')[:12],
            f"{emp['days_present']}/{emp['working_days']}",
            format_inr(emp['earned_basic']), format_inr(emp['earned_hra']), format_inr(emp['earned_da']),
            format_inr(emp['earned_other_allowances']), format_inr(emp['gross_salary']),
            format_inr(emp['pf_employee']), format_inr(emp['esi_employee']),
            format_inr(emp['professional_tax']), format_inr(emp['tds']),
            format_inr(emp['total_deductions']), format_inr(emp['net_pay']),
        ])

    totals = ['', '', f'TOTAL ({len(payroll.get("employees", []))})', '', '',
              '', '', '', '', format_inr(payroll['total_gross']),
              '', '', '', '', format_inr(payroll['total_deductions']), format_inr(payroll['total_net_pay'])]
    data.append(totals)

    col_widths = [22, 38, 68, 50, 30, 38, 35, 35, 35, 42, 35, 32, 28, 32, 42, 45]
    table = Table(data, colWidths=col_widths)
    table.setStyle(TableStyle([
        ('FONTNAME', (0, 0), (-1, -1), 'Helvetica'),
        ('FONTSIZE', (0, 0), (-1, -1), 6.5),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 7),
        ('BACKGROUND', (0, 0), (-1, 0), colors.Color(0.12, 0.35, 0.55)),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ALIGN', (4, 0), (-1, -1), 'RIGHT'),
        ('ALIGN', (0, 0), (0, -1), 'CENTER'),
        ('GRID', (0, 0), (-1, -1), 0.4, colors.lightgrey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -2), [colors.white, colors.Color(0.97, 0.97, 0.97)]),
        ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
        ('BACKGROUND', (0, -1), (-1, -1), colors.Color(0.88, 0.93, 0.88)),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
    ]))
    elements.append(table)
    doc.build(elements)
    buf.seek(0)
    year = payroll['year']
    return Response(content=buf.read(), media_type="application/pdf",
                    headers={"Content-Disposition": f'attachment; filename="Payroll_{month_name}_{year}.pdf"'})


# ============ PAYROLL REPORT (EXCEL) ============

@router.get("/hrms/reports/payroll/{payroll_id}/excel")
async def export_payroll_report_excel(payroll_id: str, user: dict = Depends(get_current_user)):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

    payroll = await db.hrms_payroll.find_one({'id': payroll_id}, {'_id': 0})
    if not payroll:
        raise HTTPException(status_code=404, detail="Payroll not found")

    settings = await db.hrms_settings.find_one({'key': 'company_info'}, {'_id': 0})
    company = get_company_info_sync(settings)
    month_name = calendar.month_name[payroll['month']]

    wb = Workbook()
    ws = wb.active
    ws.title = "Payroll"

    header_font = Font(name='Arial', size=14, bold=True)
    sub_font = Font(name='Arial', size=9, color='666666')
    col_font = Font(name='Arial', size=9, bold=True, color='FFFFFF')
    cell_font = Font(name='Arial', size=8)
    num_font = Font(name='Arial', size=8)
    bold_font = Font(name='Arial', size=9, bold=True)
    header_fill = PatternFill(start_color='1E5A8C', end_color='1E5A8C', fill_type='solid')
    total_fill = PatternFill(start_color='E0EDDF', end_color='E0EDDF', fill_type='solid')
    thin_border = Border(left=Side(style='thin', color='CCCCCC'), right=Side(style='thin', color='CCCCCC'),
                         top=Side(style='thin', color='CCCCCC'), bottom=Side(style='thin', color='CCCCCC'))

    cols = 16
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=cols)
    ws['A1'] = company['name']
    ws['A1'].font = header_font
    ws['A1'].alignment = Alignment(horizontal='center')

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=cols)
    ws['A2'] = f"Payroll Register - {month_name} {payroll['year']} | Status: {payroll['status'].upper()}"
    ws['A2'].font = sub_font
    ws['A2'].alignment = Alignment(horizontal='center')

    headers = ['S.No', 'Emp ID', 'Name', 'Dept', 'Days', 'Basic', 'HRA', 'DA', 'Other', 'Gross', 'PF', 'ESI', 'PT', 'TDS', 'Deductions', 'Net Pay']
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_idx, value=h)
        cell.font = col_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
        cell.border = thin_border

    for i, emp in enumerate(payroll.get('employees', [])):
        row = i + 5
        values = [i + 1, emp.get('employee_code', ''), emp['name'], emp.get('department', ''),
                  f"{emp['days_present']}/{emp['working_days']}",
                  emp['earned_basic'], emp['earned_hra'], emp['earned_da'], emp['earned_other_allowances'],
                  emp['gross_salary'], emp['pf_employee'], emp['esi_employee'], emp['professional_tax'],
                  emp['tds'], emp['total_deductions'], emp['net_pay']]
        for col_idx, val in enumerate(values, 1):
            cell = ws.cell(row=row, column=col_idx, value=val)
            cell.font = cell_font if col_idx <= 5 else num_font
            cell.border = thin_border
            if col_idx >= 6:
                cell.number_format = '#,##0.00'

    total_row = len(payroll.get('employees', [])) + 5
    ws.cell(row=total_row, column=3, value='TOTAL').font = bold_font
    ws.cell(row=total_row, column=10, value=payroll['total_gross']).font = bold_font
    ws.cell(row=total_row, column=15, value=payroll['total_deductions']).font = bold_font
    ws.cell(row=total_row, column=16, value=payroll['total_net_pay']).font = bold_font
    for c in range(1, cols + 1):
        cell = ws.cell(row=total_row, column=c)
        cell.fill = total_fill
        cell.border = thin_border
        if c >= 6:
            cell.number_format = '#,##0.00'

    widths = [5, 8, 20, 14, 6, 8, 8, 7, 7, 10, 7, 7, 6, 7, 10, 10]
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    year = payroll['year']
    return Response(content=buf.read(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="Payroll_{month_name}_{year}.xlsx"'})


# ============ ATTENDANCE REPORT (EXCEL) ============

@router.get("/hrms/reports/attendance/excel")
async def export_attendance_excel(month: str, department_id: Optional[str] = None, user: dict = Depends(get_current_user)):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

    settings = await db.hrms_settings.find_one({'key': 'company_info'}, {'_id': 0})
    company = get_company_info_sync(settings)
    parts = month.split('-')
    year_num, month_num = int(parts[0]), int(parts[1])
    month_name = calendar.month_name[month_num]

    emp_query = {'is_active': True}
    if department_id and department_id != 'all':
        emp_query['department_id'] = department_id

    summaries = []
    async for emp in db.hrms_employees.find(emp_query, {'_id': 0}).sort('name', 1):
        dept = await db.hrms_departments.find_one({'id': emp.get('department_id')}, {'_id': 0, 'name': 1})
        present = await db.hrms_attendance.count_documents({'employee_id': emp['id'], 'date': {'$regex': f'^{month}'}, 'status': 'present'})
        absent = await db.hrms_attendance.count_documents({'employee_id': emp['id'], 'date': {'$regex': f'^{month}'}, 'status': 'absent'})
        half_day = await db.hrms_attendance.count_documents({'employee_id': emp['id'], 'date': {'$regex': f'^{month}'}, 'status': 'half_day'})
        late = await db.hrms_attendance.count_documents({'employee_id': emp['id'], 'date': {'$regex': f'^{month}'}, 'status': 'late'})
        leave = await db.hrms_attendance.count_documents({'employee_id': emp['id'], 'date': {'$regex': f'^{month}'}, 'status': 'leave'})
        effective = present + late + (half_day * 0.5)
        summaries.append({
            'code': emp.get('employee_id', ''), 'name': emp['name'],
            'dept': dept['name'] if dept else 'N/A',
            'present': present, 'absent': absent, 'half_day': half_day,
            'late': late, 'leave': leave, 'effective': effective,
        })

    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance"

    header_font = Font(name='Arial', size=14, bold=True)
    sub_font = Font(name='Arial', size=9, color='666666')
    col_font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
    cell_font = Font(name='Arial', size=9)
    bold_font = Font(name='Arial', size=10, bold=True)
    header_fill = PatternFill(start_color='264785', end_color='264785', fill_type='solid')
    total_fill = PatternFill(start_color='E5EAF0', end_color='E5EAF0', fill_type='solid')
    thin_border = Border(left=Side(style='thin', color='CCCCCC'), right=Side(style='thin', color='CCCCCC'),
                         top=Side(style='thin', color='CCCCCC'), bottom=Side(style='thin', color='CCCCCC'))

    ws.merge_cells('A1:I1')
    ws['A1'] = company['name']
    ws['A1'].font = header_font
    ws['A1'].alignment = Alignment(horizontal='center')

    ws.merge_cells('A2:I2')
    ws['A2'] = f"Attendance Report - {month_name} {year_num}"
    ws['A2'].font = sub_font
    ws['A2'].alignment = Alignment(horizontal='center')

    headers = ['S.No', 'Emp ID', 'Name', 'Department', 'Present', 'Absent', 'Half Day', 'Late', 'Leave', 'Effective']
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col, value=h)
        cell.font = col_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
        cell.border = thin_border

    for i, s in enumerate(summaries):
        row = i + 5
        values = [i + 1, s['code'], s['name'], s['dept'], s['present'], s['absent'], s['half_day'], s['late'], s['leave'], s['effective']]
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=val)
            cell.font = cell_font
            cell.border = thin_border
            if col >= 5:
                cell.alignment = Alignment(horizontal='center')

    total_row = len(summaries) + 5
    ws.cell(row=total_row, column=3, value='TOTAL').font = bold_font
    for col, key in [(5, 'present'), (6, 'absent'), (7, 'half_day'), (8, 'late'), (9, 'leave'), (10, 'effective')]:
        val = sum(s[key] for s in summaries)
        cell = ws.cell(row=total_row, column=col, value=val)
        cell.font = bold_font
        cell.fill = total_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal='center')

    for col in [('A', 6), ('B', 10), ('C', 22), ('D', 16), ('E', 10), ('F', 10), ('G', 10), ('H', 8), ('I', 8), ('J', 10)]:
        ws.column_dimensions[col[0]].width = col[1]

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return Response(content=buf.read(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="Attendance_{month_name}_{year_num}.xlsx"'})
