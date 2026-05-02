from fastapi import APIRouter, HTTPException, Depends, UploadFile, File, Response
from database import db
from deps import get_current_user
from helpers import log_audit
from datetime import datetime, timezone
from typing import Optional
from io import BytesIO
import uuid
import base64
import re


def _safe_float(val, default=0):
    """Safely convert value to float, returning default for empty/invalid values."""
    if val is None or val == '':
        return default
    try:
        return float(val)
    except (ValueError, TypeError):
        return default


router = APIRouter()


BULK_TEMPLATE_HEADERS = [
    'Employee_ID', 'Full_Name', 'Gender', 'Date_of_Birth', 'Phone', 'Email',
    'Address', 'Department', 'Designation', 'Date_of_Joining', 'Employment_Type',
    'Basic_Salary', 'HRA', 'Allowances', 'Bank_Account_Number', 'IFSC_Code',
    'PAN_Number', 'Aadhaar_Number', 'PF_Applicable', 'ESI_Applicable'
]


def _parse_date_ddmmyyyy(val):
    """Parse DD-MM-YYYY or DD/MM/YYYY to YYYY-MM-DD"""
    if not val:
        return ''
    s = str(val).strip()
    for fmt in ('%d-%m-%Y', '%d/%m/%Y', '%Y-%m-%d'):
        try:
            return datetime.strptime(s, fmt).strftime('%Y-%m-%d')
        except ValueError:
            continue
    return None


def _validate_pan(pan):
    if not pan:
        return True
    return bool(re.match(r'^[A-Z]{5}[0-9]{4}[A-Z]$', str(pan).strip().upper()))


def _validate_aadhaar(aadhaar):
    if not aadhaar:
        return True
    cleaned = re.sub(r'\s', '', str(aadhaar).strip())
    return bool(re.match(r'^\d{12}$', cleaned))


def _capitalize_name(name):
    if not name:
        return ''
    return ' '.join(w.capitalize() for w in str(name).strip().split())


# ============ EMPLOYEES CRUD ============

@router.get("/hrms/employees")
async def get_employees(
    search: Optional[str] = None,
    department_id: Optional[str] = None,
    status: Optional[str] = None,
    page: int = 1,
    limit: int = 20,
    user: dict = Depends(get_current_user)
):
    query = {}
    if search:
        query['$or'] = [
            {'name': {'$regex': search, '$options': 'i'}},
            {'employee_id': {'$regex': search, '$options': 'i'}},
            {'email': {'$regex': search, '$options': 'i'}},
            {'phone': {'$regex': search, '$options': 'i'}},
        ]
    if department_id and department_id != 'all':
        query['department_id'] = department_id
    if status and status != 'all':
        query['is_active'] = status == 'active'

    total = await db.hrms_employees.count_documents(query)
    skip = (page - 1) * limit
    employees = []
    async for emp in db.hrms_employees.find(query, {'_id': 0}).sort('name', 1).skip(skip).limit(limit):
        dept = await db.hrms_departments.find_one({'id': emp.get('department_id')}, {'_id': 0})
        emp['department_name'] = dept['name'] if dept else 'Unassigned'
        employees.append(emp)

    return {
        'employees': employees,
        'total': total,
        'page': page,
        'total_pages': max(1, (total + limit - 1) // limit)
    }


# Upload logs endpoint - must be before {employee_id} route to avoid path conflict
@router.get("/hrms/employees/upload-logs")
async def get_upload_logs(user: dict = Depends(get_current_user)):
    if user.get('role') not in ('admin', 'hr_admin'):
        raise HTTPException(status_code=403, detail="Admin or HR Admin required")
    logs = []
    async for log in db.hrms_upload_logs.find({}, {'_id': 0}).sort('created_at', -1).limit(20):
        logs.append(log)
    return logs


@router.get("/hrms/employees/{employee_id}")
async def get_employee(employee_id: str, user: dict = Depends(get_current_user)):
    emp = await db.hrms_employees.find_one({'id': employee_id}, {'_id': 0})
    if not emp:
        raise HTTPException(status_code=404, detail="Employee not found")
    dept = await db.hrms_departments.find_one({'id': emp.get('department_id')}, {'_id': 0})
    emp['department_name'] = dept['name'] if dept else 'Unassigned'
    return emp


@router.post("/hrms/employees")
async def create_employee(data: dict, user: dict = Depends(get_current_user)):
    if user.get('role') not in ('admin', 'hr_admin'):
        raise HTTPException(status_code=403, detail="Admin or HR Admin required")

    required = ['name', 'email', 'phone', 'department_id', 'designation', 'date_of_joining']
    for field in required:
        if not data.get(field):
            raise HTTPException(status_code=400, detail=f"{field} is required")

    existing = await db.hrms_employees.find_one({'email': data['email']})
    if existing:
        raise HTTPException(status_code=400, detail="Employee with this email already exists")

    # Generate employee ID
    count = await db.hrms_employees.count_documents({})
    emp_id_num = f"K3-{count + 1:04d}"

    employee = {
        'id': str(uuid.uuid4()),
        'employee_id': emp_id_num,
        'name': data['name'],
        'email': data['email'],
        'phone': data['phone'],
        'date_of_birth': data.get('date_of_birth', ''),
        'gender': data.get('gender', ''),
        'address': data.get('address', ''),
        'department_id': data['department_id'],
        'designation': data['designation'],
        'date_of_joining': data['date_of_joining'],
        'employment_type': data.get('employment_type', 'full_time'),
        'basic_salary': _safe_float(data.get('basic_salary')),
        'hra': _safe_float(data.get('hra')),
        'da': _safe_float(data.get('da')),
        'other_allowances': _safe_float(data.get('other_allowances')),
        'pf_number': data.get('pf_number', ''),
        'esi_number': data.get('esi_number', ''),
        'pan_number': data.get('pan_number', ''),
        'aadhar_number': data.get('aadhar_number', ''),
        'bank_name': data.get('bank_name', ''),
        'bank_account_no': data.get('bank_account_no', ''),
        'ifsc_code': data.get('ifsc_code', ''),
        'photo_url': data.get('photo_url', ''),
        'emergency_contact_name': data.get('emergency_contact_name', ''),
        'emergency_contact_phone': data.get('emergency_contact_phone', ''),
        'is_active': True,
        'status': 'active',
        'increment_history': [],
        'created_by': user.get('id', ''),
        'created_at': datetime.now(timezone.utc).isoformat(),
        'updated_at': datetime.now(timezone.utc).isoformat(),
    }
    await db.hrms_employees.insert_one(employee)
    await log_audit(user.get('id', ''), user.get('name', ''), 'create', 'employee', employee['id'], f"Created employee: {employee['name']}")

    dept = await db.hrms_departments.find_one({'id': employee['department_id']}, {'_id': 0})
    employee['department_name'] = dept['name'] if dept else 'Unassigned'
    return {k: v for k, v in employee.items() if k != '_id'}


@router.put("/hrms/employees/{employee_id}")
async def update_employee(employee_id: str, data: dict, user: dict = Depends(get_current_user)):
    if user.get('role') not in ('admin', 'hr_admin'):
        raise HTTPException(status_code=403, detail="Admin or HR Admin required")
    emp = await db.hrms_employees.find_one({'id': employee_id})
    if not emp:
        raise HTTPException(status_code=404, detail="Employee not found")

    updatable = [
        'name', 'email', 'phone', 'date_of_birth', 'gender', 'address',
        'department_id', 'designation', 'date_of_joining', 'employment_type',
        'basic_salary', 'hra', 'da', 'other_allowances',
        'pf_number', 'esi_number', 'pan_number', 'aadhar_number',
        'bank_name', 'bank_account_no', 'ifsc_code',
        'emergency_contact_name', 'emergency_contact_phone',
        'is_active', 'status', 'photo_url', 'tax_regime'
    ]
    update_fields = {}
    for field in updatable:
        if field in data:
            if field in ('basic_salary', 'hra', 'da', 'other_allowances'):
                update_fields[field] = _safe_float(data[field])
            else:
                update_fields[field] = data[field]

    if not update_fields:
        raise HTTPException(status_code=400, detail="No valid fields to update")

    update_fields['updated_at'] = datetime.now(timezone.utc).isoformat()
    await db.hrms_employees.update_one({'id': employee_id}, {'$set': update_fields})
    await log_audit(user.get('id', ''), user.get('name', ''), 'update', 'employee', employee_id, f"Updated employee fields: {list(update_fields.keys())}")

    updated = await db.hrms_employees.find_one({'id': employee_id}, {'_id': 0})
    dept = await db.hrms_departments.find_one({'id': updated.get('department_id')}, {'_id': 0})
    updated['department_name'] = dept['name'] if dept else 'Unassigned'
    return updated


@router.delete("/hrms/employees/{employee_id}")
async def delete_employee(employee_id: str, user: dict = Depends(get_current_user)):
    if user.get('role') not in ('admin', 'hr_admin'):
        raise HTTPException(status_code=403, detail="Admin or HR Admin required")
    result = await db.hrms_employees.update_one(
        {'id': employee_id},
        {'$set': {'is_active': False, 'status': 'inactive', 'updated_at': datetime.now(timezone.utc).isoformat()}}
    )
    if result.modified_count == 0:
        raise HTTPException(status_code=404, detail="Employee not found")
    await log_audit(user.get('id', ''), user.get('name', ''), 'delete', 'employee', employee_id, "Deactivated employee")
    return {"message": "Employee deactivated"}


@router.post("/hrms/employees/{employee_id}/photo")
async def upload_employee_photo(employee_id: str, file: UploadFile = File(...), user: dict = Depends(get_current_user)):
    if user.get('role') not in ('admin', 'hr_admin'):
        raise HTTPException(status_code=403, detail="Admin or HR Admin required")
    emp = await db.hrms_employees.find_one({'id': employee_id})
    if not emp:
        raise HTTPException(status_code=404, detail="Employee not found")
    content = await file.read()
    if len(content) > 2 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="File must be under 2MB")
    encoded = base64.b64encode(content).decode('utf-8')
    content_type = file.content_type or 'image/png'
    data_uri = f"data:{content_type};base64,{encoded}"
    await db.hrms_employees.update_one(
        {'id': employee_id},
        {'$set': {'photo_url': data_uri, 'updated_at': datetime.now(timezone.utc).isoformat()}}
    )
    return {"message": "Photo uploaded", "photo_url": data_uri}


# ============ INCREMENT HISTORY ============

@router.post("/hrms/employees/{employee_id}/increment")
async def add_increment(employee_id: str, data: dict, user: dict = Depends(get_current_user)):
    if user.get('role') not in ('admin', 'hr_admin'):
        raise HTTPException(status_code=403, detail="Admin or HR Admin required")
    emp = await db.hrms_employees.find_one({'id': employee_id})
    if not emp:
        raise HTTPException(status_code=404, detail="Employee not found")

    increment = {
        'id': str(uuid.uuid4()),
        'date': data.get('date', datetime.now(timezone.utc).strftime('%Y-%m-%d')),
        'previous_salary': emp.get('basic_salary', 0),
        'new_salary': _safe_float(data.get('new_salary')),
        'increment_amount': _safe_float(data.get('new_salary')) - emp.get('basic_salary', 0),
        'reason': data.get('reason', ''),
        'approved_by': user.get('name', ''),
        'created_at': datetime.now(timezone.utc).isoformat()
    }

    await db.hrms_employees.update_one(
        {'id': employee_id},
        {
            '$push': {'increment_history': increment},
            '$set': {
                'basic_salary': _safe_float(data.get('new_salary')),
                'updated_at': datetime.now(timezone.utc).isoformat()
            }
        }
    )
    return {"message": "Increment added", "increment": increment}


# ============ EMPLOYEE BULK UPLOAD ============

@router.get("/hrms/employees/bulk-upload/template")
async def download_bulk_template(user: dict = Depends(get_current_user)):
    if user.get('role') not in ('admin', 'hr_admin'):
        raise HTTPException(status_code=403, detail="Admin or HR Admin required")
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        raise HTTPException(status_code=500, detail="openpyxl not installed")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Employee Data"

    header_font = Font(name='Arial', bold=True, size=11, color='FFFFFF')
    header_fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
    header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    col_widths = {
        'A': 14, 'B': 22, 'C': 10, 'D': 14, 'E': 14, 'F': 25,
        'G': 30, 'H': 18, 'I': 18, 'J': 14, 'K': 16,
        'L': 14, 'M': 10, 'N': 12, 'O': 22, 'P': 14,
        'Q': 14, 'R': 16, 'S': 14, 'T': 14
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    for col_idx, header in enumerate(BULK_TEMPLATE_HEADERS, 1):
        cell = ws.cell(row=1, column=col_idx, value=header.replace('_', ' '))
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    # Instructions sheet
    ws2 = wb.create_sheet("Instructions")
    instructions = [
        ["K3 GAS SERVICE - Employee Bulk Upload Instructions"],
        [""],
        ["Field", "Required", "Format", "Example"],
        ["Employee_ID", "Optional", "Auto-generated if blank", "K3-0001"],
        ["Full_Name", "YES", "Text", "Rajesh Kumar"],
        ["Gender", "YES", "male/female/other", "male"],
        ["Date_of_Birth", "Optional", "DD-MM-YYYY", "15-06-1990"],
        ["Phone", "YES", "10 digits", "9876543210"],
        ["Email", "YES", "Valid email", "rajesh@email.com"],
        ["Address", "Optional", "Text", "123 Main Street"],
        ["Department", "YES", "Exact department name", "Administration"],
        ["Designation", "YES", "Text", "Manager"],
        ["Date_of_Joining", "YES", "DD-MM-YYYY", "01-01-2024"],
        ["Employment_Type", "Optional", "full_time/part_time/contract", "full_time"],
        ["Basic_Salary", "Optional", "Number", "25000"],
        ["HRA", "Optional", "Number", "5000"],
        ["Allowances", "Optional", "Number", "3000"],
        ["Bank_Account_Number", "Optional", "Text", "1234567890"],
        ["IFSC_Code", "Optional", "Text", "SBIN0001234"],
        ["PAN_Number", "Optional", "ABCDE1234F", "ABCDE1234F"],
        ["Aadhaar_Number", "Optional", "12 digits", "123456789012"],
        ["PF_Applicable", "Optional", "yes/no", "yes"],
        ["ESI_Applicable", "Optional", "yes/no", "no"],
    ]

    inst_font = Font(name='Arial', size=11)
    title_font = Font(name='Arial', size=14, bold=True, color='1F4E79')
    for row_idx, row_data in enumerate(instructions, 1):
        for col_idx, value in enumerate(row_data, 1):
            cell = ws2.cell(row=row_idx, column=col_idx, value=value)
            cell.font = title_font if row_idx == 1 else inst_font
    ws2.column_dimensions['A'].width = 22
    ws2.column_dimensions['B'].width = 12
    ws2.column_dimensions['C'].width = 28
    ws2.column_dimensions['D'].width = 18

    # Available departments sheet
    ws3 = wb.create_sheet("Departments")
    ws3.cell(row=1, column=1, value="Department Name").font = Font(name='Arial', bold=True, size=11)
    ws3.column_dimensions['A'].width = 30
    row = 2
    async for dept in db.hrms_departments.find({}, {'_id': 0, 'name': 1}).sort('name', 1):
        ws3.cell(row=row, column=1, value=dept['name'])
        row += 1

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=employee_bulk_upload_template.xlsx"}
    )


@router.post("/hrms/employees/bulk-upload/validate")
async def validate_bulk_upload(file: UploadFile = File(...), user: dict = Depends(get_current_user)):
    if user.get('role') not in ('admin', 'hr_admin'):
        raise HTTPException(status_code=403, detail="Admin or HR Admin required")
    try:
        import openpyxl
    except ImportError:
        raise HTTPException(status_code=500, detail="openpyxl not installed")

    content = await file.read()
    if len(content) > 5 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="File must be under 5MB")

    try:
        wb = openpyxl.load_workbook(BytesIO(content), data_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="Invalid Excel file")

    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    if not rows:
        raise HTTPException(status_code=400, detail="No data rows found in the file")

    # Build department lookup
    dept_lookup = {}
    async for dept in db.hrms_departments.find({}, {'_id': 0, 'id': 1, 'name': 1}):
        dept_lookup[dept['name'].lower()] = dept['id']

    # Build existing email set
    existing_emails = set()
    async for emp in db.hrms_employees.find({}, {'_id': 0, 'email': 1}):
        existing_emails.add(emp['email'].lower())

    # Build existing employee_id set
    existing_emp_ids = set()
    async for emp in db.hrms_employees.find({}, {'_id': 0, 'employee_id': 1}):
        existing_emp_ids.add(emp.get('employee_id', '').upper())

    validated = []
    errors = []
    seen_emails = set()
    seen_emp_ids = set()

    for row_idx, row in enumerate(rows, 2):
        if all(v is None for v in row):
            continue

        row_data = list(row) + [None] * (20 - len(row))
        row_errors = []

        emp_id_raw = str(row_data[0] or '').strip()
        name = _capitalize_name(row_data[1])
        gender = str(row_data[2] or '').strip().lower()
        dob_raw = str(row_data[3] or '').strip()
        phone = str(row_data[4] or '').strip()
        email_raw = str(row_data[5] or '').strip().lower()
        address = str(row_data[6] or '').strip()
        dept_name = str(row_data[7] or '').strip()
        designation = str(row_data[8] or '').strip()
        doj_raw = str(row_data[9] or '').strip()
        emp_type = str(row_data[10] or 'full_time').strip().lower()
        basic_salary_raw = row_data[11]
        hra_raw = row_data[12]
        allowances_raw = row_data[13]
        bank_acc = str(row_data[14] or '').strip()
        ifsc = str(row_data[15] or '').strip()
        pan_raw = str(row_data[16] or '').strip()
        aadhaar_raw = str(row_data[17] or '').strip()

        # Required field checks
        if not name:
            row_errors.append("Full Name is required")
        if not email_raw:
            row_errors.append("Email is required")
        if not phone:
            row_errors.append("Phone is required")
        if not dept_name:
            row_errors.append("Department is required")
        if not designation:
            row_errors.append("Designation is required")
        if not doj_raw:
            row_errors.append("Date of Joining is required")
        if not gender or gender not in ('male', 'female', 'other'):
            row_errors.append("Gender must be male/female/other")

        # Date validation
        dob_parsed = _parse_date_ddmmyyyy(dob_raw) if dob_raw else ''
        if dob_raw and dob_parsed is None:
            row_errors.append("Invalid Date of Birth format (use DD-MM-YYYY)")

        doj_parsed = _parse_date_ddmmyyyy(doj_raw) if doj_raw else ''
        if doj_raw and doj_parsed is None:
            row_errors.append("Invalid Date of Joining format (use DD-MM-YYYY)")

        # Department check
        dept_id = dept_lookup.get(dept_name.lower(), None) if dept_name else None
        if dept_name and not dept_id:
            row_errors.append(f"Department '{dept_name}' not found. Check 'Departments' sheet for valid names")

        # Duplicate email checks
        if email_raw:
            if email_raw in seen_emails:
                row_errors.append(f"Duplicate email in file: {email_raw}")
            if email_raw in existing_emails:
                row_errors.append(f"Email already exists in database: {email_raw}")
            seen_emails.add(email_raw)

        # Employee ID duplicate
        if emp_id_raw:
            emp_id_upper = emp_id_raw.upper()
            if emp_id_upper in seen_emp_ids:
                row_errors.append(f"Duplicate Employee ID in file: {emp_id_raw}")
            if emp_id_upper in existing_emp_ids:
                row_errors.append(f"Employee ID already exists: {emp_id_raw}")
            seen_emp_ids.add(emp_id_upper)

        # PAN & Aadhaar validation
        if pan_raw and not _validate_pan(pan_raw):
            row_errors.append(f"Invalid PAN format: {pan_raw} (expected: ABCDE1234F)")
        if aadhaar_raw and not _validate_aadhaar(aadhaar_raw):
            row_errors.append(f"Invalid Aadhaar format: {aadhaar_raw} (expected: 12 digits)")

        # Salary parsing
        try:
            basic_salary = float(basic_salary_raw) if basic_salary_raw else 0
        except (ValueError, TypeError):
            basic_salary = 0
            row_errors.append("Invalid Basic Salary (must be a number)")
        try:
            hra = float(hra_raw) if hra_raw else 0
        except (ValueError, TypeError):
            hra = 0
            row_errors.append("Invalid HRA (must be a number)")
        try:
            allowances = float(allowances_raw) if allowances_raw else 0
        except (ValueError, TypeError):
            allowances = 0
            row_errors.append("Invalid Allowances (must be a number)")

        entry = {
            'row': row_idx,
            'employee_id': emp_id_raw,
            'name': name,
            'gender': gender,
            'date_of_birth': dob_parsed or '',
            'phone': phone,
            'email': email_raw,
            'address': address,
            'department': dept_name,
            'department_id': dept_id or '',
            'designation': designation,
            'date_of_joining': doj_parsed or '',
            'employment_type': emp_type if emp_type in ('full_time', 'part_time', 'contract') else 'full_time',
            'basic_salary': basic_salary,
            'hra': hra,
            'other_allowances': allowances,
            'bank_account_no': bank_acc,
            'ifsc_code': ifsc,
            'pan_number': pan_raw.upper() if pan_raw else '',
            'aadhar_number': re.sub(r'\s', '', aadhaar_raw) if aadhaar_raw else '',
            'errors': row_errors,
            'status': 'error' if row_errors else 'valid',
        }
        validated.append(entry)
        if row_errors:
            errors.append({'row': row_idx, 'errors': row_errors})

    valid_count = sum(1 for v in validated if v['status'] == 'valid')
    error_count = sum(1 for v in validated if v['status'] == 'error')

    return {
        'total_rows': len(validated),
        'valid_count': valid_count,
        'error_count': error_count,
        'data': validated,
        'errors': errors,
        'available_departments': list(dept_lookup.keys()),
    }


@router.post("/hrms/employees/bulk-upload/confirm")
async def confirm_bulk_upload(data: dict, user: dict = Depends(get_current_user)):
    if user.get('role') not in ('admin', 'hr_admin'):
        raise HTTPException(status_code=403, detail="Admin or HR Admin required")

    employees_data = data.get('employees', [])
    mode = data.get('mode', 'skip')

    if not employees_data:
        raise HTTPException(status_code=400, detail="No employee data provided")

    created = 0
    skipped = 0
    updated = 0
    failed = []

    # Get next employee ID number
    max_emp = await db.hrms_employees.count_documents({})
    next_id = max_emp + 1

    for emp in employees_data:
        if emp.get('status') == 'error' and mode == 'skip':
            skipped += 1
            continue

        try:
            email = emp.get('email', '').lower()
            existing = await db.hrms_employees.find_one({'email': email})

            if existing:
                if mode == 'overwrite':
                    update_fields = {
                        'name': emp.get('name', existing.get('name')),
                        'gender': emp.get('gender', existing.get('gender')),
                        'phone': emp.get('phone', existing.get('phone')),
                        'designation': emp.get('designation', existing.get('designation')),
                        'department_id': emp.get('department_id', existing.get('department_id')),
                        'basic_salary': _safe_float(emp.get('basic_salary', existing.get('basic_salary', 0))),
                        'hra': _safe_float(emp.get('hra', existing.get('hra', 0))),
                        'other_allowances': _safe_float(emp.get('other_allowances', existing.get('other_allowances', 0))),
                        'updated_at': datetime.now(timezone.utc).isoformat(),
                    }
                    if emp.get('date_of_birth'):
                        update_fields['date_of_birth'] = emp['date_of_birth']
                    if emp.get('date_of_joining'):
                        update_fields['date_of_joining'] = emp['date_of_joining']
                    if emp.get('address'):
                        update_fields['address'] = emp['address']
                    if emp.get('bank_account_no'):
                        update_fields['bank_account_no'] = emp['bank_account_no']
                    if emp.get('ifsc_code'):
                        update_fields['ifsc_code'] = emp['ifsc_code']
                    if emp.get('pan_number'):
                        update_fields['pan_number'] = emp['pan_number']
                    if emp.get('aadhar_number'):
                        update_fields['aadhar_number'] = emp['aadhar_number']

                    await db.hrms_employees.update_one({'email': email}, {'$set': update_fields})
                    updated += 1
                else:
                    skipped += 1
                continue

            emp_id = emp.get('employee_id') or f"K3-{next_id:04d}"
            next_id += 1

            new_emp = {
                'id': str(uuid.uuid4()),
                'employee_id': emp_id,
                'name': emp.get('name', ''),
                'email': email,
                'phone': emp.get('phone', ''),
                'date_of_birth': emp.get('date_of_birth', ''),
                'gender': emp.get('gender', ''),
                'address': emp.get('address', ''),
                'department_id': emp.get('department_id', ''),
                'designation': emp.get('designation', ''),
                'date_of_joining': emp.get('date_of_joining', ''),
                'employment_type': emp.get('employment_type', 'full_time'),
                'basic_salary': _safe_float(emp.get('basic_salary')),
                'hra': _safe_float(emp.get('hra')),
                'da': 0,
                'other_allowances': _safe_float(emp.get('other_allowances')),
                'pf_number': '',
                'esi_number': '',
                'pan_number': emp.get('pan_number', ''),
                'aadhar_number': emp.get('aadhar_number', ''),
                'bank_name': '',
                'bank_account_no': emp.get('bank_account_no', ''),
                'ifsc_code': emp.get('ifsc_code', ''),
                'photo_url': '',
                'emergency_contact_name': '',
                'emergency_contact_phone': '',
                'is_active': True,
                'status': 'active',
                'increment_history': [],
                'created_by': user.get('id', ''),
                'created_at': datetime.now(timezone.utc).isoformat(),
                'updated_at': datetime.now(timezone.utc).isoformat(),
            }
            await db.hrms_employees.insert_one(new_emp)
            created += 1
        except Exception as e:
            failed.append({'email': emp.get('email', 'unknown'), 'error': str(e)})

    # Log the upload
    log_entry = {
        'id': str(uuid.uuid4()),
        'uploaded_by': user.get('name', ''),
        'uploaded_by_id': user.get('id', ''),
        'total_rows': len(employees_data),
        'created': created,
        'updated': updated,
        'skipped': skipped,
        'failed': len(failed),
        'mode': mode,
        'created_at': datetime.now(timezone.utc).isoformat(),
    }
    await db.hrms_upload_logs.insert_one(log_entry)

    return {
        'message': f"Bulk upload complete: {created} created, {updated} updated, {skipped} skipped, {len(failed)} failed",
        'created': created,
        'updated': updated,
        'skipped': skipped,
        'failed': failed,
    }


@router.post("/hrms/employees/bulk-upload/error-report")
async def download_error_report(data: dict, user: dict = Depends(get_current_user)):
    if user.get('role') not in ('admin', 'hr_admin'):
        raise HTTPException(status_code=403, detail="Admin or HR Admin required")
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        raise HTTPException(status_code=500, detail="openpyxl not installed")

    errors = data.get('errors', [])
    if not errors:
        raise HTTPException(status_code=400, detail="No errors to report")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Error Report"

    header_font = Font(name='Arial', bold=True, size=11, color='FFFFFF')
    header_fill = PatternFill(start_color='C00000', end_color='C00000', fill_type='solid')
    error_fill = PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid')
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    headers_list = ['Row', 'Name', 'Email', 'Phone', 'Department', 'Errors']
    for col_idx, header in enumerate(headers_list, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal='center')

    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 22
    ws.column_dimensions['C'].width = 25
    ws.column_dimensions['D'].width = 14
    ws.column_dimensions['E'].width = 18
    ws.column_dimensions['F'].width = 60

    for row_idx, err_entry in enumerate(errors, 2):
        row_data = err_entry if isinstance(err_entry, dict) else {}
        ws.cell(row=row_idx, column=1, value=row_data.get('row', row_idx)).border = thin_border
        ws.cell(row=row_idx, column=2, value=row_data.get('name', '')).border = thin_border
        ws.cell(row=row_idx, column=3, value=row_data.get('email', '')).border = thin_border
        ws.cell(row=row_idx, column=4, value=row_data.get('phone', '')).border = thin_border
        ws.cell(row=row_idx, column=5, value=row_data.get('department', '')).border = thin_border
        error_list = row_data.get('errors', [])
        error_text = '; '.join(error_list) if isinstance(error_list, list) else str(error_list)
        cell = ws.cell(row=row_idx, column=6, value=error_text)
        cell.border = thin_border
        cell.fill = error_fill

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=bulk_upload_error_report.xlsx"}
    )
