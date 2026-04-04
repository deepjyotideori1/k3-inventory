from fastapi import APIRouter, HTTPException, Depends, UploadFile, File, Response
from database import db
from deps import get_current_user
from helpers import log_audit
from datetime import datetime, timezone
from typing import Optional
from io import BytesIO
import uuid
import calendar
import re

router = APIRouter()


# ============ LEAVE CONFIGURATION ============

DEFAULT_LEAVE_TYPES = [
    {'id': 'casual', 'name': 'Casual Leave', 'annual_quota': 12, 'carry_forward': False},
    {'id': 'sick', 'name': 'Sick Leave', 'annual_quota': 10, 'carry_forward': False},
    {'id': 'earned', 'name': 'Earned Leave', 'annual_quota': 15, 'carry_forward': True},
    {'id': 'unpaid', 'name': 'Unpaid Leave', 'annual_quota': 0, 'carry_forward': False},
]


@router.get("/hrms/leave/config")
async def get_leave_config(user: dict = Depends(get_current_user)):
    config = await db.hrms_leave_config.find_one({'key': 'leave_types'}, {'_id': 0})
    if not config:
        doc = {'key': 'leave_types', 'leave_types': DEFAULT_LEAVE_TYPES}
        await db.hrms_leave_config.insert_one(doc)
        return {'leave_types': DEFAULT_LEAVE_TYPES}
    return {'leave_types': config.get('leave_types', [])}


@router.put("/hrms/leave/config")
async def update_leave_config(data: dict, user: dict = Depends(get_current_user)):
    if user.get('role') not in ('admin', 'hr_admin'):
        raise HTTPException(status_code=403, detail="Admin or HR Admin required")
    leave_types = data.get('leave_types', [])
    await db.hrms_leave_config.update_one(
        {'key': 'leave_types'},
        {'$set': {'leave_types': leave_types}},
        upsert=True
    )
    return {"message": "Leave configuration updated"}


# ============ ATTENDANCE MARKING ============

@router.post("/hrms/attendance/mark")
async def mark_attendance(data: dict, user: dict = Depends(get_current_user)):
    if user.get('role') not in ('admin', 'hr_admin'):
        raise HTTPException(status_code=403, detail="Admin or HR Admin required")

    records = data.get('records', [])
    date = data.get('date', '')
    if not date:
        raise HTTPException(status_code=400, detail="Date is required")
    if not records:
        raise HTTPException(status_code=400, detail="At least one attendance record is required")

    created = 0
    updated = 0
    for rec in records:
        emp_id = rec.get('employee_id')
        status = rec.get('status', 'present')
        if status not in ('present', 'absent', 'half_day', 'late', 'leave', 'holiday', 'week_off'):
            continue

        existing = await db.hrms_attendance.find_one({'employee_id': emp_id, 'date': date})
        if existing:
            await db.hrms_attendance.update_one(
                {'employee_id': emp_id, 'date': date},
                {'$set': {
                    'status': status,
                    'leave_type': rec.get('leave_type', ''),
                    'check_in': rec.get('check_in', ''),
                    'check_out': rec.get('check_out', ''),
                    'overtime_hours': float(rec.get('overtime_hours', 0)),
                    'remarks': rec.get('remarks', ''),
                    'updated_by': user.get('name', ''),
                    'updated_at': datetime.now(timezone.utc).isoformat(),
                }}
            )
            updated += 1
        else:
            att = {
                'id': str(uuid.uuid4()),
                'employee_id': emp_id,
                'date': date,
                'status': status,
                'leave_type': rec.get('leave_type', ''),
                'check_in': rec.get('check_in', ''),
                'check_out': rec.get('check_out', ''),
                'overtime_hours': float(rec.get('overtime_hours', 0)),
                'remarks': rec.get('remarks', ''),
                'marked_by': user.get('name', ''),
                'created_at': datetime.now(timezone.utc).isoformat(),
            }
            await db.hrms_attendance.insert_one(att)
            created += 1

    return {"message": f"Attendance marked: {created} new, {updated} updated", "created": created, "updated": updated}


@router.get("/hrms/attendance")
async def get_attendance(
    date: Optional[str] = None,
    employee_id: Optional[str] = None,
    month: Optional[str] = None,
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    page: int = 1,
    limit: int = 50,
    user: dict = Depends(get_current_user)
):
    query = {}
    # Employee role: restrict to own records
    if user.get('role') == 'hrms_employee':
        linked = user.get('linked_employee_id', '')
        if linked:
            query['employee_id'] = linked
        else:
            return {'records': [], 'total': 0, 'page': 1, 'total_pages': 1}
    if date:
        query['date'] = date
    if employee_id and user.get('role') != 'hrms_employee':
        query['employee_id'] = employee_id
    if start_date and end_date:
        query['date'] = {'$gte': start_date, '$lte': end_date}
    elif month:
        query['date'] = {'$regex': f'^{month}'}

    total = await db.hrms_attendance.count_documents(query)
    skip = (page - 1) * limit
    records = []
    async for rec in db.hrms_attendance.find(query, {'_id': 0}).sort('date', 1).skip(skip).limit(limit):
        emp = await db.hrms_employees.find_one({'id': rec['employee_id']}, {'_id': 0, 'name': 1, 'employee_id': 1, 'department_id': 1, 'designation': 1})
        if emp:
            rec['employee_name'] = emp.get('name', '')
            rec['employee_code'] = emp.get('employee_id', '')
            dept = await db.hrms_departments.find_one({'id': emp.get('department_id')}, {'_id': 0, 'name': 1})
            rec['department'] = dept['name'] if dept else 'Unassigned'
        records.append(rec)

    return {
        'records': records,
        'total': total,
        'page': page,
        'total_pages': max(1, (total + limit - 1) // limit)
    }


@router.put("/hrms/attendance/{record_id}")
async def update_attendance(record_id: str, data: dict, user: dict = Depends(get_current_user)):
    if user.get('role') not in ('admin', 'hr_admin'):
        raise HTTPException(status_code=403, detail="Admin or HR Admin required")
    allowed = ['status', 'leave_type', 'check_in', 'check_out', 'overtime_hours', 'remarks']
    update_fields = {k: v for k, v in data.items() if k in allowed}
    if 'overtime_hours' in update_fields:
        update_fields['overtime_hours'] = float(update_fields['overtime_hours'])
    update_fields['updated_by'] = user.get('name', '')
    update_fields['updated_at'] = datetime.now(timezone.utc).isoformat()

    result = await db.hrms_attendance.update_one({'id': record_id}, {'$set': update_fields})
    if result.modified_count == 0:
        raise HTTPException(status_code=404, detail="Attendance record not found")
    return {"message": "Attendance updated"}


# ============ MONTHLY ATTENDANCE SUMMARY ============

@router.get("/hrms/attendance/monthly-summary")
async def get_monthly_attendance_summary(
    month: str,
    department_id: Optional[str] = None,
    user: dict = Depends(get_current_user)
):
    """month format: YYYY-MM"""
    emp_query = {'is_active': True}
    if department_id and department_id != 'all':
        emp_query['department_id'] = department_id

    parts = month.split('-')
    year_num = int(parts[0])
    month_num = int(parts[1])
    _, total_days = calendar.monthrange(year_num, month_num)

    summaries = []
    async for emp in db.hrms_employees.find(emp_query, {'_id': 0}).sort('name', 1):
        dept = await db.hrms_departments.find_one({'id': emp.get('department_id')}, {'_id': 0, 'name': 1})

        present = await db.hrms_attendance.count_documents({
            'employee_id': emp['id'], 'date': {'$regex': f'^{month}'}, 'status': 'present'
        })
        absent = await db.hrms_attendance.count_documents({
            'employee_id': emp['id'], 'date': {'$regex': f'^{month}'}, 'status': 'absent'
        })
        half_day = await db.hrms_attendance.count_documents({
            'employee_id': emp['id'], 'date': {'$regex': f'^{month}'}, 'status': 'half_day'
        })
        late = await db.hrms_attendance.count_documents({
            'employee_id': emp['id'], 'date': {'$regex': f'^{month}'}, 'status': 'late'
        })
        leave = await db.hrms_attendance.count_documents({
            'employee_id': emp['id'], 'date': {'$regex': f'^{month}'}, 'status': 'leave'
        })
        holiday = await db.hrms_attendance.count_documents({
            'employee_id': emp['id'], 'date': {'$regex': f'^{month}'}, 'status': 'holiday'
        })
        week_off = await db.hrms_attendance.count_documents({
            'employee_id': emp['id'], 'date': {'$regex': f'^{month}'}, 'status': 'week_off'
        })

        pipeline = [
            {'$match': {'employee_id': emp['id'], 'date': {'$regex': f'^{month}'}}},
            {'$group': {'_id': None, 'total_ot': {'$sum': '$overtime_hours'}}}
        ]
        ot_result = await db.hrms_attendance.aggregate(pipeline).to_list(1)
        total_ot = ot_result[0]['total_ot'] if ot_result else 0

        effective_present = present + late + (half_day * 0.5)

        summaries.append({
            'employee_id': emp['id'],
            'employee_code': emp.get('employee_id', ''),
            'name': emp['name'],
            'department': dept['name'] if dept else 'Unassigned',
            'designation': emp.get('designation', ''),
            'total_days': total_days,
            'present': present,
            'absent': absent,
            'half_day': half_day,
            'late': late,
            'leave': leave,
            'holiday': holiday,
            'week_off': week_off,
            'overtime_hours': round(total_ot, 1),
            'effective_present': effective_present,
        })

    return {
        'month': month,
        'total_days': total_days,
        'summaries': summaries,
    }


# ============ LEAVE BALANCE ============

@router.get("/hrms/leave/balance/{employee_id}")
async def get_leave_balance(employee_id: str, year: Optional[int] = None, user: dict = Depends(get_current_user)):
    if not year:
        year = datetime.now().year

    emp = await db.hrms_employees.find_one({'id': employee_id}, {'_id': 0})
    if not emp:
        raise HTTPException(status_code=404, detail="Employee not found")

    config = await db.hrms_leave_config.find_one({'key': 'leave_types'}, {'_id': 0})
    leave_types = config.get('leave_types', DEFAULT_LEAVE_TYPES) if config else DEFAULT_LEAVE_TYPES

    year_str = str(year)
    balances = []
    for lt in leave_types:
        taken = await db.hrms_attendance.count_documents({
            'employee_id': employee_id,
            'date': {'$regex': f'^{year_str}'},
            'status': 'leave',
            'leave_type': lt['id']
        })
        balances.append({
            'leave_type_id': lt['id'],
            'leave_type_name': lt['name'],
            'annual_quota': lt['annual_quota'],
            'taken': taken,
            'remaining': max(0, lt['annual_quota'] - taken) if lt['annual_quota'] > 0 else 'N/A',
        })

    return {
        'employee_id': employee_id,
        'employee_name': emp.get('name', ''),
        'year': year,
        'balances': balances
    }


# ============ DAILY ATTENDANCE LIST FOR A DATE ============

@router.get("/hrms/attendance/daily")
async def get_daily_attendance(
    date: str,
    department_id: Optional[str] = None,
    user: dict = Depends(get_current_user)
):
    """Get all employees with their attendance for a specific date"""
    emp_query = {'is_active': True}
    if department_id and department_id != 'all':
        emp_query['department_id'] = department_id

    result = []
    async for emp in db.hrms_employees.find(emp_query, {'_id': 0}).sort('name', 1):
        dept = await db.hrms_departments.find_one({'id': emp.get('department_id')}, {'_id': 0, 'name': 1})
        att = await db.hrms_attendance.find_one({'employee_id': emp['id'], 'date': date}, {'_id': 0})
        result.append({
            'employee_id': emp['id'],
            'employee_code': emp.get('employee_id', ''),
            'name': emp['name'],
            'department': dept['name'] if dept else 'Unassigned',
            'designation': emp.get('designation', ''),
            'status': att.get('status', '') if att else '',
            'check_in': att.get('check_in', '') if att else '',
            'check_out': att.get('check_out', '') if att else '',
            'overtime_hours': att.get('overtime_hours', 0) if att else 0,
            'leave_type': att.get('leave_type', '') if att else '',
            'remarks': att.get('remarks', '') if att else '',
            'record_id': att.get('id', '') if att else '',
        })

    return {'date': date, 'employees': result}


# ============ EMPLOYEE-WISE ATTENDANCE OVERVIEW (Custom Date Range) ============

@router.get("/hrms/attendance/employee-overview")
async def get_employee_attendance_overview(
    employee_id: str,
    start_date: str,
    end_date: str,
    user: dict = Depends(get_current_user)
):
    """Get detailed attendance overview for a single employee across a custom date range"""
    emp = await db.hrms_employees.find_one({'id': employee_id}, {'_id': 0})
    if not emp:
        raise HTTPException(status_code=404, detail="Employee not found")

    dept = await db.hrms_departments.find_one({'id': emp.get('department_id')}, {'_id': 0, 'name': 1})

    query = {'employee_id': employee_id, 'date': {'$gte': start_date, '$lte': end_date}}

    records = []
    async for rec in db.hrms_attendance.find(query, {'_id': 0}).sort('date', 1):
        records.append(rec)

    # Compute summary stats
    present = sum(1 for r in records if r.get('status') == 'present')
    absent = sum(1 for r in records if r.get('status') == 'absent')
    half_day = sum(1 for r in records if r.get('status') == 'half_day')
    late = sum(1 for r in records if r.get('status') == 'late')
    leave = sum(1 for r in records if r.get('status') == 'leave')
    holiday = sum(1 for r in records if r.get('status') == 'holiday')
    week_off = sum(1 for r in records if r.get('status') == 'week_off')
    total_ot = sum(float(r.get('overtime_hours', 0)) for r in records)
    effective_present = present + late + (half_day * 0.5)

    # Leave type breakdown
    leave_breakdown = {}
    for r in records:
        if r.get('status') == 'leave' and r.get('leave_type'):
            lt = r['leave_type']
            leave_breakdown[lt] = leave_breakdown.get(lt, 0) + 1

    return {
        'employee': {
            'id': emp['id'],
            'name': emp['name'],
            'employee_code': emp.get('employee_id', ''),
            'department': dept['name'] if dept else 'Unassigned',
            'designation': emp.get('designation', ''),
        },
        'date_range': {'start': start_date, 'end': end_date},
        'summary': {
            'total_records': len(records),
            'present': present,
            'absent': absent,
            'half_day': half_day,
            'late': late,
            'leave': leave,
            'holiday': holiday,
            'week_off': week_off,
            'overtime_hours': round(total_ot, 1),
            'effective_present': effective_present,
            'leave_breakdown': leave_breakdown,
        },
        'records': records,
    }


# ============ BIOMETRIC SYNC ENDPOINT (API-Ready) ============

@router.post("/hrms/attendance/biometric-sync")
async def biometric_sync(data: dict, user: dict = Depends(get_current_user)):
    """
    API-ready endpoint for biometric device integration.
    Accepts bulk attendance punches from external biometric systems.
    Expected payload:
    {
        "device_id": "BIO-001",
        "punches": [
            {"employee_code": "K3-001", "timestamp": "2026-04-15T09:05:00", "type": "check_in"},
            {"employee_code": "K3-001", "timestamp": "2026-04-15T18:10:00", "type": "check_out"}
        ]
    }
    """
    if user.get('role') not in ('admin', 'hr_admin'):
        raise HTTPException(status_code=403, detail="Admin or HR Admin required")

    device_id = data.get('device_id', 'unknown')
    punches = data.get('punches', [])
    if not punches:
        raise HTTPException(status_code=400, detail="No punches provided")

    processed = 0
    errors = []
    for punch in punches:
        emp_code = punch.get('employee_code', '')
        timestamp_str = punch.get('timestamp', '')
        punch_type = punch.get('type', '')  # check_in or check_out

        if not emp_code or not timestamp_str or punch_type not in ('check_in', 'check_out'):
            errors.append(f"Invalid punch data: {punch}")
            continue

        emp = await db.hrms_employees.find_one({'employee_id': emp_code}, {'_id': 0, 'id': 1})
        if not emp:
            errors.append(f"Employee not found: {emp_code}")
            continue

        try:
            ts = datetime.fromisoformat(timestamp_str)
            punch_date = ts.strftime('%Y-%m-%d')
            punch_time = ts.strftime('%H:%M')
        except ValueError:
            errors.append(f"Invalid timestamp: {timestamp_str}")
            continue

        existing = await db.hrms_attendance.find_one({'employee_id': emp['id'], 'date': punch_date})
        if existing:
            update_field = {punch_type: punch_time, 'updated_at': datetime.now(timezone.utc).isoformat(), 'updated_by': f'biometric:{device_id}'}
            if punch_type == 'check_in' and not existing.get('status'):
                update_field['status'] = 'present'
            await db.hrms_attendance.update_one({'employee_id': emp['id'], 'date': punch_date}, {'$set': update_field})
        else:
            att = {
                'id': str(uuid.uuid4()),
                'employee_id': emp['id'],
                'date': punch_date,
                'status': 'present' if punch_type == 'check_in' else '',
                'leave_type': '',
                'check_in': punch_time if punch_type == 'check_in' else '',
                'check_out': punch_time if punch_type == 'check_out' else '',
                'overtime_hours': 0,
                'remarks': f'Biometric ({device_id})',
                'marked_by': f'biometric:{device_id}',
                'created_at': datetime.now(timezone.utc).isoformat(),
            }
            await db.hrms_attendance.insert_one(att)
        processed += 1

    return {
        "message": f"Biometric sync complete: {processed} punches processed, {len(errors)} errors",
        "processed": processed,
        "errors": errors,
    }


# ============ BULK ATTENDANCE UPLOAD ============

ATTENDANCE_TEMPLATE_HEADERS = [
    'Employee_ID', 'Date', 'Status', 'Check_In', 'Check_Out', 'OT_Hours', 'Leave_Type', 'Remarks'
]

VALID_STATUSES = {'present', 'absent', 'half_day', 'late', 'leave', 'holiday', 'week_off'}
VALID_LEAVE_TYPES = {'casual', 'sick', 'earned', 'unpaid'}


def _parse_att_date(val):
    if not val:
        return None
    s = str(val).strip()
    for fmt in ('%d-%m-%Y', '%d/%m/%Y', '%Y-%m-%d'):
        try:
            return datetime.strptime(s, fmt).strftime('%Y-%m-%d')
        except ValueError:
            continue
    return None


def _parse_time(val):
    if not val:
        return ''
    s = str(val).strip()
    # Handle HH:MM or HH:MM:SS
    if re.match(r'^\d{1,2}:\d{2}(:\d{2})?$', s):
        return s[:5]
    return ''


@router.get("/hrms/attendance/bulk-upload/template")
async def download_attendance_template(user: dict = Depends(get_current_user)):
    if user.get('role') not in ('admin', 'hr_admin'):
        raise HTTPException(status_code=403, detail="Admin or HR Admin required")
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance Upload"

    hf = Font(name='Arial', size=10, bold=True, color='FFFFFF')
    hfill = PatternFill(start_color='264785', end_color='264785', fill_type='solid')
    border = Border(
        left=Side(style='thin', color='CCCCCC'), right=Side(style='thin', color='CCCCCC'),
        top=Side(style='thin', color='CCCCCC'), bottom=Side(style='thin', color='CCCCCC')
    )

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(ATTENDANCE_TEMPLATE_HEADERS))
    ws['A1'] = 'Attendance Bulk Upload Template'
    ws['A1'].font = Font(name='Arial', size=14, bold=True)
    ws['A1'].alignment = Alignment(horizontal='center')

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(ATTENDANCE_TEMPLATE_HEADERS))
    ws['A2'] = 'Required: Employee_ID, Date, Status. Date format: DD-MM-YYYY. One row per employee per date.'
    ws['A2'].font = Font(name='Arial', size=8, color='666666')
    ws['A2'].alignment = Alignment(horizontal='center')

    for col, h in enumerate(ATTENDANCE_TEMPLATE_HEADERS, 1):
        cell = ws.cell(row=4, column=col, value=h)
        cell.font = hf
        cell.fill = hfill
        cell.alignment = Alignment(horizontal='center')
        cell.border = border

    # Sample rows
    samples = [
        ['K3-0003', '01-04-2026', 'Present', '09:00', '18:00', 0, '', 'On time'],
        ['K3-0003', '02-04-2026', 'Late', '09:35', '18:00', 0, '', 'Traffic'],
        ['K3-0003', '03-04-2026', 'Leave', '', '', 0, 'Casual', 'Family function'],
        ['K3-0007', '01-04-2026', 'Present', '08:55', '18:30', 0.5, '', ''],
        ['K3-0007', '02-04-2026', 'Absent', '', '', 0, '', 'No call'],
    ]
    for r, row_data in enumerate(samples, 5):
        for c, val in enumerate(row_data, 1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.font = Font(name='Arial', size=9, color='888888', italic=True)
            cell.border = border

    # Data validations
    status_dv = DataValidation(type="list", formula1='"Present,Absent,Half Day,Late,Leave,Holiday,Week Off"', allow_blank=False)
    status_dv.error = "Invalid status"
    status_dv.errorTitle = "Status"
    ws.add_data_validation(status_dv)
    status_dv.add('C5:C500')

    leave_dv = DataValidation(type="list", formula1='"Casual,Sick,Earned,Unpaid"', allow_blank=True)
    ws.add_data_validation(leave_dv)
    leave_dv.add('G5:G500')

    widths = [14, 14, 12, 10, 10, 10, 12, 24]
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w

    # Instructions sheet
    instr = wb.create_sheet("Instructions")
    instructions = [
        ['Field', 'Required', 'Format / Notes'],
        ['Employee_ID', 'Yes', 'Employee code as in system (e.g., K3-0003)'],
        ['Date', 'Yes', 'DD-MM-YYYY format. One row per employee per date.'],
        ['Status', 'Yes', 'Present / Absent / Half Day / Late / Leave / Holiday / Week Off'],
        ['Check_In', 'No', 'HH:MM (24hr) e.g., 09:00'],
        ['Check_Out', 'No', 'HH:MM (24hr) e.g., 18:00'],
        ['OT_Hours', 'No', 'Overtime hours (numeric, e.g., 1.5)'],
        ['Leave_Type', 'If Leave', 'Casual / Sick / Earned / Unpaid (required when Status = Leave)'],
        ['Remarks', 'No', 'Optional notes'],
    ]
    for r, row in enumerate(instructions, 1):
        for c, val in enumerate(row, 1):
            cell = instr.cell(row=r, column=c, value=val)
            if r == 1:
                cell.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
                cell.fill = PatternFill(start_color='264785', end_color='264785', fill_type='solid')
            else:
                cell.font = Font(name='Arial', size=9)
    instr.column_dimensions['A'].width = 16
    instr.column_dimensions['B'].width = 10
    instr.column_dimensions['C'].width = 55

    # Employee list sheet for reference
    emp_sheet = wb.create_sheet("Employee List")
    emp_sheet.cell(row=1, column=1, value='Employee_ID').font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
    emp_sheet.cell(row=1, column=2, value='Name').font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
    emp_sheet.cell(row=1, column=3, value='Department').font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
    for c in range(1, 4):
        emp_sheet.cell(row=1, column=c).fill = PatternFill(start_color='264785', end_color='264785', fill_type='solid')

    row_idx = 2
    async for emp in db.hrms_employees.find({'is_active': True}, {'_id': 0, 'employee_id': 1, 'name': 1, 'department_id': 1}).sort('employee_id', 1):
        emp_sheet.cell(row=row_idx, column=1, value=emp.get('employee_id', '')).font = Font(name='Arial', size=9)
        emp_sheet.cell(row=row_idx, column=2, value=emp.get('name', '')).font = Font(name='Arial', size=9)
        dept = await db.hrms_departments.find_one({'id': emp.get('department_id')}, {'_id': 0, 'name': 1})
        emp_sheet.cell(row=row_idx, column=3, value=dept['name'] if dept else '').font = Font(name='Arial', size=9)
        row_idx += 1
    emp_sheet.column_dimensions['A'].width = 14
    emp_sheet.column_dimensions['B'].width = 22
    emp_sheet.column_dimensions['C'].width = 18

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return Response(content=buf.read(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": 'attachment; filename="Attendance_Upload_Template.xlsx"'})


@router.post("/hrms/attendance/bulk-upload/validate")
async def validate_attendance_upload(file: UploadFile = File(...), user: dict = Depends(get_current_user)):
    if user.get('role') not in ('admin', 'hr_admin'):
        raise HTTPException(status_code=403, detail="Admin or HR Admin required")
    if not file.filename.endswith('.xlsx'):
        raise HTTPException(status_code=400, detail="Only .xlsx files are accepted")

    from openpyxl import load_workbook

    content = await file.read()
    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="File too large (max 10MB)")

    try:
        wb = load_workbook(BytesIO(content), data_only=True)
        ws = wb.active
    except Exception:
        raise HTTPException(status_code=400, detail="Invalid Excel file")

    # Find header row
    header_row = None
    for row_idx in range(1, min(10, ws.max_row + 1)):
        vals = [str(ws.cell(row=row_idx, column=c).value or '').strip() for c in range(1, ws.max_column + 1)]
        if 'Employee_ID' in vals and 'Date' in vals:
            header_row = row_idx
            break
    if not header_row:
        raise HTTPException(status_code=400, detail="Could not find headers. Ensure 'Employee_ID' and 'Date' columns exist.")

    headers = [str(ws.cell(row=header_row, column=c).value or '').strip() for c in range(1, ws.max_column + 1)]
    col_map = {}
    for i, h in enumerate(headers):
        for tmpl in ATTENDANCE_TEMPLATE_HEADERS:
            if h.lower().replace(' ', '_') == tmpl.lower() or h.lower() == tmpl.lower():
                col_map[tmpl] = i
                break

    # Load employee map
    emp_map = {}
    async for emp in db.hrms_employees.find({'is_active': True}, {'_id': 0, 'id': 1, 'employee_id': 1, 'name': 1}):
        emp_map[emp.get('employee_id', '').strip().upper()] = {'id': emp['id'], 'name': emp.get('name', ''), 'code': emp.get('employee_id', '')}

    STATUS_MAP = {
        'present': 'present', 'absent': 'absent', 'half_day': 'half_day', 'half day': 'half_day',
        'late': 'late', 'leave': 'leave', 'holiday': 'holiday', 'week_off': 'week_off', 'week off': 'week_off',
    }
    LEAVE_MAP = {
        'casual': 'casual', 'casual leave': 'casual', 'sick': 'sick', 'sick leave': 'sick',
        'earned': 'earned', 'earned leave': 'earned', 'unpaid': 'unpaid', 'unpaid leave': 'unpaid',
    }

    validated_rows = []
    errors = []
    seen_keys = set()

    for row_idx in range(header_row + 1, ws.max_row + 1):
        row_vals = [ws.cell(row=row_idx, column=c + 1).value for c in range(len(headers))]
        if all(v is None or str(v).strip() == '' for v in row_vals):
            continue

        def get_val(col_name):
            idx = col_map.get(col_name)
            if idx is None:
                return ''
            val = row_vals[idx]
            return str(val).strip() if val is not None else ''

        row_num = row_idx
        row_errors = []

        emp_code = get_val('Employee_ID').strip().upper()
        date_raw = get_val('Date')
        status_raw = get_val('Status').strip().lower()

        # Required checks
        if not emp_code:
            row_errors.append('Employee_ID is required')
        if not date_raw:
            row_errors.append('Date is required')
        if not status_raw:
            row_errors.append('Status is required')

        # Validate employee
        emp_info = emp_map.get(emp_code)
        if emp_code and not emp_info:
            row_errors.append(f'Employee not found: {emp_code}')

        # Parse date
        parsed_date = _parse_att_date(date_raw) if date_raw else None
        if date_raw and not parsed_date:
            row_errors.append(f'Invalid date format: {date_raw} (use DD-MM-YYYY)')

        # Validate status
        status = STATUS_MAP.get(status_raw, '')
        if status_raw and not status:
            row_errors.append(f'Invalid status: {status_raw}. Use: Present/Absent/Late/Leave/Half Day/Holiday/Week Off')

        # Check duplicate (same employee + same date)
        if emp_code and parsed_date:
            key = f"{emp_code}|{parsed_date}"
            if key in seen_keys:
                row_errors.append(f'Duplicate entry in file: {emp_code} on {parsed_date}')
            seen_keys.add(key)

        # Check if exists in DB
        is_existing = False
        if emp_info and parsed_date and not row_errors:
            existing = await db.hrms_attendance.find_one({'employee_id': emp_info['id'], 'date': parsed_date})
            is_existing = existing is not None

        # Leave type
        leave_raw = get_val('Leave_Type').strip().lower()
        leave_type = LEAVE_MAP.get(leave_raw, '')
        if status == 'leave' and not leave_type:
            row_errors.append('Leave_Type is required when Status is Leave')

        # Parse times
        check_in = _parse_time(get_val('Check_In'))
        check_out = _parse_time(get_val('Check_Out'))

        # OT hours
        ot_raw = get_val('OT_Hours')
        try:
            ot_hours = float(ot_raw) if ot_raw else 0
        except (ValueError, TypeError):
            row_errors.append(f'Invalid OT_Hours: {ot_raw}')
            ot_hours = 0

        row_data = {
            'row_num': row_num,
            'employee_code': emp_info['code'] if emp_info else emp_code,
            'employee_id': emp_info['id'] if emp_info else '',
            'employee_name': emp_info['name'] if emp_info else '',
            'date': parsed_date or date_raw,
            'status': status,
            'check_in': check_in,
            'check_out': check_out,
            'overtime_hours': ot_hours,
            'leave_type': leave_type,
            'remarks': get_val('Remarks'),
            'is_existing': is_existing,
            'errors': row_errors,
            'has_errors': len(row_errors) > 0,
        }
        validated_rows.append(row_data)
        if row_errors:
            errors.append({'row': row_num, 'employee_code': emp_code, 'date': date_raw, 'errors': row_errors})

    total = len(validated_rows)
    valid_count = sum(1 for r in validated_rows if not r['has_errors'])
    error_count = sum(1 for r in validated_rows if r['has_errors'])
    new_count = sum(1 for r in validated_rows if not r['has_errors'] and not r['is_existing'])
    update_count = sum(1 for r in validated_rows if not r['has_errors'] and r['is_existing'])

    return {
        'total_rows': total,
        'valid_count': valid_count,
        'error_count': error_count,
        'new_count': new_count,
        'update_count': update_count,
        'rows': validated_rows,
        'errors': errors,
    }


@router.post("/hrms/attendance/bulk-upload/confirm")
async def confirm_attendance_upload(data: dict, user: dict = Depends(get_current_user)):
    if user.get('role') not in ('admin', 'hr_admin'):
        raise HTTPException(status_code=403, detail="Admin or HR Admin required")

    rows = data.get('rows', [])
    mode = data.get('mode', 'skip')

    if not rows:
        raise HTTPException(status_code=400, detail="No rows to import")

    created = 0
    updated = 0
    skipped = 0

    for row in rows:
        if row.get('has_errors') or not row.get('employee_id') or not row.get('date'):
            skipped += 1
            continue

        existing = await db.hrms_attendance.find_one({'employee_id': row['employee_id'], 'date': row['date']})

        if existing:
            if mode == 'skip':
                skipped += 1
                continue
            await db.hrms_attendance.update_one(
                {'employee_id': row['employee_id'], 'date': row['date']},
                {'$set': {
                    'status': row.get('status', ''),
                    'check_in': row.get('check_in', ''),
                    'check_out': row.get('check_out', ''),
                    'overtime_hours': float(row.get('overtime_hours', 0)),
                    'leave_type': row.get('leave_type', ''),
                    'remarks': row.get('remarks', ''),
                    'updated_at': datetime.now(timezone.utc).isoformat(),
                    'marked_by': f'bulk_upload:{user.get("name", "")}',
                }}
            )
            updated += 1
        else:
            att = {
                'id': str(uuid.uuid4()),
                'employee_id': row['employee_id'],
                'date': row['date'],
                'status': row.get('status', ''),
                'leave_type': row.get('leave_type', ''),
                'check_in': row.get('check_in', ''),
                'check_out': row.get('check_out', ''),
                'overtime_hours': float(row.get('overtime_hours', 0)),
                'remarks': row.get('remarks', ''),
                'marked_by': f'bulk_upload:{user.get("name", "")}',
                'created_at': datetime.now(timezone.utc).isoformat(),
            }
            await db.hrms_attendance.insert_one(att)
            created += 1

    # Log upload
    log = {
        'id': str(uuid.uuid4()),
        'type': 'attendance',
        'uploaded_by': user.get('name', ''),
        'uploaded_by_id': user.get('id', ''),
        'mode': mode,
        'total_rows': len(rows),
        'created': created,
        'updated': updated,
        'skipped': skipped,
        'created_at': datetime.now(timezone.utc).isoformat(),
    }
    await db.hrms_upload_logs.insert_one(log)
    await log_audit(user.get('id', ''), user.get('name', ''), 'bulk_upload', 'attendance', log['id'],
                    f"Attendance bulk upload: {created} created, {updated} updated, {skipped} skipped")

    return {
        "message": f"Import complete: {created} created, {updated} updated, {skipped} skipped",
        "created": created,
        "updated": updated,
        "skipped": skipped,
    }


@router.post("/hrms/attendance/bulk-upload/error-report")
async def download_attendance_error_report(data: dict, user: dict = Depends(get_current_user)):
    if user.get('role') not in ('admin', 'hr_admin'):
        raise HTTPException(status_code=403, detail="Admin or HR Admin required")
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

    errors = data.get('errors', [])
    wb = Workbook()
    ws = wb.active
    ws.title = "Error Report"

    hf = Font(name='Arial', size=10, bold=True, color='FFFFFF')
    ef = Font(name='Arial', size=9, color='CC0000')
    nf = Font(name='Arial', size=9)
    hfill = PatternFill(start_color='CC3333', end_color='CC3333', fill_type='solid')
    border = Border(left=Side(style='thin', color='CCCCCC'), right=Side(style='thin', color='CCCCCC'),
                    top=Side(style='thin', color='CCCCCC'), bottom=Side(style='thin', color='CCCCCC'))

    ws.merge_cells('A1:D1')
    ws['A1'] = 'Attendance Bulk Upload - Error Report'
    ws['A1'].font = Font(name='Arial', size=14, bold=True, color='CC0000')
    ws['A1'].alignment = Alignment(horizontal='center')

    headers = ['Row #', 'Employee ID', 'Date', 'Errors']
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col, value=h)
        cell.font = hf
        cell.fill = hfill
        cell.border = border

    for i, err in enumerate(errors):
        row = i + 4
        ws.cell(row=row, column=1, value=err.get('row', '')).font = nf
        ws.cell(row=row, column=2, value=err.get('employee_code', '')).font = nf
        ws.cell(row=row, column=3, value=err.get('date', '')).font = nf
        ws.cell(row=row, column=4, value='; '.join(err.get('errors', []))).font = ef
        for c in range(1, 5):
            ws.cell(row=row, column=c).border = border

    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 16
    ws.column_dimensions['C'].width = 14
    ws.column_dimensions['D'].width = 60

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return Response(content=buf.read(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": 'attachment; filename="Attendance_Upload_Error_Report.xlsx"'})
