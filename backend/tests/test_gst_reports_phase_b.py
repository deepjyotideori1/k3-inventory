"""
GST Billing Phase B - Reports module tests.
Tests 10 report builders + Excel/PDF exports + admin-only access.
"""
import os
import io
import pytest
import requests
from openpyxl import load_workbook

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "").rstrip("/")
API = f"{BASE_URL}/api"

ADMIN = {"email": "admin@k3gas.com", "password": "Admin@123"}
NON_ADMIN = {"email": "jullang@k3gas.com", "password": "Jullang@123"}

REPORT_KEYS = [
    "daily_sales", "monthly_sales", "gst_report", "hsn_summary", "item_wise",
    "customer_wise", "warehouse_wise", "cancelled", "payment_wise", "tax_summary",
]


def _login(creds):
    r = requests.post(f"{API}/auth/login", json=creds, timeout=30)
    assert r.status_code == 200, f"login failed: {r.status_code} {r.text}"
    d = r.json()
    return d.get("token") or d.get("access_token")


@pytest.fixture(scope="module")
def admin_headers():
    return {"Authorization": f"Bearer {_login(ADMIN)}"}


@pytest.fixture(scope="module")
def non_admin_headers():
    return {"Authorization": f"Bearer {_login(NON_ADMIN)}"}


# ---------- LIST ----------
class TestReportList:
    def test_list_returns_10_reports(self, admin_headers):
        r = requests.get(f"{API}/gst/reports/list", headers=admin_headers, timeout=30)
        assert r.status_code == 200
        data = r.json()
        assert isinstance(data, list)
        assert len(data) == 10
        keys = {item["key"] for item in data}
        assert keys == set(REPORT_KEYS)
        for item in data:
            assert item.get("label")


# ---------- DATA ENDPOINTS ----------
class TestReportData:
    def _get(self, key, admin_headers, params=""):
        r = requests.get(f"{API}/gst/reports/{key}{params}", headers=admin_headers, timeout=60)
        assert r.status_code == 200, f"{key} failed: {r.status_code} {r.text[:300]}"
        d = r.json()
        assert d["report_type"] == key
        assert "label" in d
        assert "columns" in d
        assert "rows" in d
        assert "summary" in d
        return d

    def test_daily_sales(self, admin_headers):
        d = self._get("daily_sales", admin_headers)
        assert len(d["columns"]) == 8
        col_keys = [c["key"] for c in d["columns"]]
        assert col_keys == ["date", "invoices", "sub_total", "cgst", "sgst", "igst", "total_gst", "grand_total"]
        for k in ["sub_total", "cgst", "sgst", "igst", "total_gst", "grand_total", "invoices"]:
            assert k in d["summary"]
        if d["rows"]:
            r0 = d["rows"][0]
            for k in ["date", "invoices", "sub_total", "cgst", "sgst", "igst", "total_gst", "grand_total"]:
                assert k in r0

    def test_monthly_sales(self, admin_headers):
        d = self._get("monthly_sales", admin_headers)
        col_keys = [c["key"] for c in d["columns"]]
        assert col_keys[0] == "month"
        if d["rows"]:
            m = d["rows"][0]["month"]
            # YYYY-MM format
            assert len(m) == 7 and m[4] == "-", f"month not YYYY-MM: {m!r}"

    def test_gst_report(self, admin_headers):
        d = self._get("gst_report", admin_headers)
        assert "b2b_count" in d["summary"]
        assert "b2c_count" in d["summary"]
        assert len(d["columns"]) == 13
        for r in d["rows"]:
            assert r["type"] in ("B2B", "B2C")
            if r["type"] == "B2B":
                assert r["customer_gstin"] and r["customer_gstin"] != "—"
            for k in ["invoice_number", "invoice_date", "customer_name", "place_of_supply",
                      "tax_mode", "sub_total", "cgst", "sgst", "igst", "total_gst", "grand_total"]:
                assert k in r

    def test_hsn_summary(self, admin_headers):
        d = self._get("hsn_summary", admin_headers)
        col_keys = [c["key"] for c in d["columns"]]
        for k in ["hsn", "unit", "gst_rate", "quantity", "taxable_value", "cgst", "sgst", "igst", "total_value"]:
            assert k in col_keys
        # Different gst_rates for same HSN are distinct rows: verify uniqueness key
        seen = set()
        for r in d["rows"]:
            key = (r["hsn"], r["unit"], r["gst_rate"])
            assert key not in seen, f"duplicate hsn|unit|gst_rate row: {key}"
            seen.add(key)

    def test_item_wise(self, admin_headers):
        d = self._get("item_wise", admin_headers)
        col_keys = [c["key"] for c in d["columns"]]
        for k in ["item_name", "hsn", "unit", "invoices", "quantity", "taxable_value", "total_gst", "total_value"]:
            assert k in col_keys
        # Sorted desc by total_value
        if len(d["rows"]) >= 2:
            for i in range(len(d["rows"]) - 1):
                assert d["rows"][i]["total_value"] >= d["rows"][i + 1]["total_value"]

    def test_customer_wise(self, admin_headers):
        d = self._get("customer_wise", admin_headers)
        col_keys = [c["key"] for c in d["columns"]]
        for k in ["customer_name", "customer_phone", "customer_gstin", "invoices",
                  "sub_total", "total_gst", "grand_total"]:
            assert k in col_keys
        if len(d["rows"]) >= 2:
            for i in range(len(d["rows"]) - 1):
                assert d["rows"][i]["grand_total"] >= d["rows"][i + 1]["grand_total"]

    def test_warehouse_wise_excludes_plant(self, admin_headers):
        d = self._get("warehouse_wise", admin_headers)
        for r in d["rows"]:
            wh = (r["warehouse"] or "").lower()
            assert "plant" not in wh, f"Plant warehouse leaked: {r['warehouse']}"
            assert "hollongi" not in wh, f"Hollongi warehouse leaked: {r['warehouse']}"

    def test_cancelled(self, admin_headers):
        d = self._get("cancelled", admin_headers)
        col_keys = [c["key"] for c in d["columns"]]
        for k in ["invoice_number", "invoice_date", "customer_name", "grand_total",
                  "cancelled_at", "cancelled_by", "cancellation_reason"]:
            assert k in col_keys
        assert "count" in d["summary"]
        assert "lost_value" in d["summary"]
        assert d["summary"]["count"] == len(d["rows"])

    def test_payment_wise(self, admin_headers):
        d = self._get("payment_wise", admin_headers)
        col_keys = [c["key"] for c in d["columns"]]
        assert col_keys[0] == "payment_mode"
        # uppercase modes
        for r in d["rows"]:
            assert r["payment_mode"] == r["payment_mode"].upper(), f"not uppercase: {r['payment_mode']}"

    def test_tax_summary(self, admin_headers):
        d = self._get("tax_summary", admin_headers)
        col_keys = [c["key"] for c in d["columns"]]
        for k in ["gst_rate", "taxable", "cgst", "sgst", "igst", "total_gst"]:
            assert k in col_keys
        # sorted asc by rate
        rates = [r["gst_rate"] for r in d["rows"]]
        assert rates == sorted(rates), f"tax_summary not sorted asc: {rates}"

    def test_invalid_report_type_returns_400(self, admin_headers):
        r = requests.get(f"{API}/gst/reports/bogus_type", headers=admin_headers, timeout=30)
        assert r.status_code == 400
        body = r.json()
        msg = str(body)
        # detail should list valid report types
        assert "daily_sales" in msg or "report_type" in msg.lower()

    def test_date_range_filter(self, admin_headers):
        # Use a narrow past range that should return no rows
        d = self._get("daily_sales", admin_headers,
                      params="?start_date=1990-01-01&end_date=1990-01-31")
        assert d["rows"] == []


# ---------- EXCEL EXPORTS ----------
class TestExcelExports:
    @pytest.mark.parametrize("key", REPORT_KEYS)
    def test_excel_export(self, admin_headers, key):
        r = requests.get(f"{API}/gst/reports/{key}/excel", headers=admin_headers, timeout=60)
        assert r.status_code == 200, f"{key}: {r.status_code} {r.text[:200]}"
        ct = r.headers.get("content-type", "")
        assert "spreadsheet" in ct or "openxml" in ct, f"{key} wrong content-type: {ct}"
        assert len(r.content) > 2000, f"{key} xlsx too small: {len(r.content)}"
        # xlsx magic = PK
        assert r.content[:2] == b"PK", f"{key} not a valid zip/xlsx"

    def test_excel_structure_daily_sales(self, admin_headers):
        r = requests.get(f"{API}/gst/reports/daily_sales/excel", headers=admin_headers, timeout=60)
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb.active
        # Sheet title equals label (truncated to 31)
        assert ws.title == "Daily Sales Report"
        # Title rows
        assert ws["A1"].value, "A1 must be company name"
        assert "GSTIN" in (ws["A2"].value or "")
        assert "Daily Sales Report" in (ws["A3"].value or "")
        # Header row at 5
        headers = [ws.cell(row=5, column=c).value for c in range(1, 9)]
        assert headers == ["Date", "Invoices", "Sub Total", "CGST", "SGST", "IGST", "Total GST", "Grand Total"]
        # Frozen at A6
        assert ws.freeze_panes == "A6", f"got {ws.freeze_panes}"
        # Autofilter starts at A5
        assert ws.auto_filter.ref and ws.auto_filter.ref.startswith("A5:")
        # Landscape since 8 columns > 6
        assert ws.page_setup.orientation == "landscape"


# ---------- PDF EXPORTS ----------
class TestPdfExports:
    @pytest.mark.parametrize("key", REPORT_KEYS)
    def test_pdf_export(self, admin_headers, key):
        r = requests.get(f"{API}/gst/reports/{key}/pdf", headers=admin_headers, timeout=60)
        assert r.status_code == 200, f"{key}: {r.status_code} {r.text[:200]}"
        ct = r.headers.get("content-type", "")
        assert "pdf" in ct, f"{key} wrong content-type: {ct}"
        assert r.content[:4] == b"%PDF", f"{key} not a PDF, starts {r.content[:8]!r}"
        assert len(r.content) > 2000, f"{key} pdf too small: {len(r.content)}"


# ---------- ADMIN-ONLY ACCESS ----------
class TestAdminOnly:
    def test_non_admin_blocked_on_list(self, non_admin_headers):
        r = requests.get(f"{API}/gst/reports/list", headers=non_admin_headers, timeout=30)
        assert r.status_code in (401, 403), f"expected 401/403, got {r.status_code}"

    def test_non_admin_blocked_on_data(self, non_admin_headers):
        r = requests.get(f"{API}/gst/reports/daily_sales", headers=non_admin_headers, timeout=30)
        assert r.status_code in (401, 403), f"expected 401/403, got {r.status_code}"

    def test_non_admin_blocked_on_excel(self, non_admin_headers):
        r = requests.get(f"{API}/gst/reports/daily_sales/excel", headers=non_admin_headers, timeout=30)
        assert r.status_code in (401, 403)

    def test_non_admin_blocked_on_pdf(self, non_admin_headers):
        r = requests.get(f"{API}/gst/reports/daily_sales/pdf", headers=non_admin_headers, timeout=30)
        assert r.status_code in (401, 403)
