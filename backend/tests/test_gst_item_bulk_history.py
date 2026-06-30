"""Backend regression for GST Item Master Bulk Update + Rate History.

Endpoints under test:
- GET  /api/gst/items/template/excel
- POST /api/gst/items/bulk-upload  (multipart, accepts .xlsx and .csv)
- POST /api/gst/items/bulk-update  (JSON)
- GET  /api/gst/items/history
- GET  /api/gst/items/history/excel
Regression touch points:
- GET /api/gst/items
- GET /api/gst/plans
- GET /api/gst/invoices?page=1&limit=20
- GET /api/gst/reports/list  (frontend uses /gst/reports/list; alt /gst-reports/list checked)
"""
import io
import os
import csv
import time
import pytest
import requests
from openpyxl import load_workbook

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "").rstrip("/")
assert BASE_URL, "REACT_APP_BACKEND_URL must be set"
API = f"{BASE_URL}/api"

ADMIN_EMAIL = "admin@k3gas.com"
ADMIN_PASS = "Admin@123"

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture(scope="session")
def admin_token():
    r = requests.post(f"{API}/auth/login",
                      json={"email": ADMIN_EMAIL, "password": ADMIN_PASS}, timeout=20)
    assert r.status_code == 200, f"login failed: {r.status_code} {r.text}"
    tok = r.json().get("token") or r.json().get("access_token")
    assert tok, f"no token in login response: {r.json()}"
    return tok


@pytest.fixture(scope="session")
def admin_headers(admin_token):
    return {"Authorization": f"Bearer {admin_token}"}


# ---------- Template Download ----------
class TestTemplateDownload:
    def test_template_returns_valid_xlsx(self, admin_headers):
        r = requests.get(f"{API}/gst/items/template/excel", headers=admin_headers, timeout=30)
        assert r.status_code == 200, r.text
        ctype = r.headers.get("content-type", "")
        assert "spreadsheetml" in ctype or "xlsx" in ctype or "octet-stream" in ctype, ctype
        assert len(r.content) > 500, "xlsx body too small"
        # Validate xlsx can be opened and has expected headers
        wb = load_workbook(io.BytesIO(r.content), data_only=True)
        ws = wb.active
        headers = [str(c.value or "").strip().lower() for c in ws[1]]
        expected = ["id", "name", "hsn", "unit", "gst_rate", "default_rate", "is_active", "effective_from"]
        for h in expected:
            assert h in headers, f"missing header {h} in {headers}"

    def test_template_requires_admin(self):
        r = requests.get(f"{API}/gst/items/template/excel", timeout=15)
        assert r.status_code in (401, 403)


# ---------- Bulk Upload via CSV ----------
class TestBulkUploadCsv:
    created_ids = []

    def test_bulk_upload_csv_creates_item(self, admin_headers):
        unique = f"TEST_BulkCsv_{int(time.time())}"
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(["id", "name", "hsn", "unit", "gst_rate", "default_rate", "is_active", "effective_from"])
        w.writerow(["", unique, "998877", "Nos", 18, 250, "active", ""])
        files = {"file": (f"{unique}.csv", buf.getvalue().encode("utf-8"), "text/csv")}
        data = {"effective_from": "2026-01-01"}
        r = requests.post(f"{API}/gst/items/bulk-upload", headers=admin_headers,
                          files=files, data=data, timeout=30)
        assert r.status_code == 200, r.text
        body = r.json()
        assert body.get("created", 0) >= 1, body
        assert body.get("rate_changes", 0) >= 1, body
        # Verify it shows up in /gst/items
        r2 = requests.get(f"{API}/gst/items", headers=admin_headers, timeout=15)
        assert r2.status_code == 200
        items = r2.json() if isinstance(r2.json(), list) else r2.json().get("items", [])
        found = next((it for it in items if it.get("name") == unique), None)
        assert found, f"Created item '{unique}' not found in /gst/items"
        assert float(found.get("default_rate", 0)) == 250.0
        assert float(found.get("gst_rate", 0)) == 18.0
        TestBulkUploadCsv.created_ids.append(found["id"])
        # Verify history entry exists for this item
        h = requests.get(f"{API}/gst/items/history",
                         params={"item_id": found["id"]}, headers=admin_headers, timeout=15)
        assert h.status_code == 200
        rows = h.json().get("rows", [])
        assert any(r.get("item_id") == found["id"] for r in rows), "no history row recorded for new item"

    def test_bulk_upload_no_file_400(self, admin_headers):
        # Missing file => FastAPI returns 422
        r = requests.post(f"{API}/gst/items/bulk-upload", headers=admin_headers, timeout=15)
        assert r.status_code in (400, 422)

    def test_bulk_upload_requires_admin(self):
        files = {"file": ("a.csv", b"name\nx", "text/csv")}
        r = requests.post(f"{API}/gst/items/bulk-upload", files=files, timeout=15)
        assert r.status_code in (401, 403)


# ---------- Bulk Update JSON (updates an existing item) ----------
class TestBulkUpdateJson:
    def test_bulk_update_existing_item_records_history(self, admin_headers):
        # First create a fresh item so we can mutate
        unique = f"TEST_BulkUpd_{int(time.time())}"
        c = requests.post(f"{API}/gst/items", headers=admin_headers,
                          json={"name": unique, "hsn": "111111", "unit": "Nos",
                                "gst_rate": 5, "default_rate": 100}, timeout=15)
        assert c.status_code in (200, 201), c.text
        new_item = c.json()
        item_id = new_item["id"]
        # Now bulk-update it via JSON: change default_rate (should create a history row)
        payload = {
            "effective_from": "2026-02-01",
            "items": [{
                "id": item_id, "name": unique, "hsn": "111111", "unit": "Nos",
                "gst_rate": 5, "default_rate": 175, "is_active": "active",
            }],
        }
        r = requests.post(f"{API}/gst/items/bulk-update", headers=admin_headers,
                          json=payload, timeout=15)
        assert r.status_code == 200, r.text
        body = r.json()
        assert body.get("updated", 0) >= 1, body
        assert body.get("rate_changes", 0) >= 1, body
        # History should contain entry with prev_rate=100 new_rate=175 effective_from=2026-02-01
        h = requests.get(f"{API}/gst/items/history",
                         params={"item_id": item_id}, headers=admin_headers, timeout=15)
        assert h.status_code == 200
        rows = h.json().get("rows", [])
        matched = [r for r in rows if float(r.get("prev_rate", -1)) == 100.0 and float(r.get("new_rate", -1)) == 175.0]
        assert matched, f"no 100->175 history row found in {rows[:3]}"
        assert matched[0]["effective_from"] == "2026-02-01"


# ---------- Bulk Upload via XLSX ----------
class TestBulkUploadXlsx:
    def test_bulk_upload_xlsx_creates_item(self, admin_headers):
        from openpyxl import Workbook
        wb = Workbook(); ws = wb.active
        ws.append(["id", "name", "hsn", "unit", "gst_rate", "default_rate", "is_active", "effective_from"])
        unique = f"TEST_BulkXlsx_{int(time.time())}"
        ws.append(["", unique, "554433", "Nos", 12, 333.5, "active", ""])
        bio = io.BytesIO(); wb.save(bio); bio.seek(0)
        files = {"file": (f"{unique}.xlsx", bio.read(), XLSX_MIME)}
        data = {"effective_from": "2026-03-15"}
        r = requests.post(f"{API}/gst/items/bulk-upload", headers=admin_headers,
                          files=files, data=data, timeout=30)
        assert r.status_code == 200, r.text
        body = r.json()
        assert body.get("created", 0) >= 1, body


# ---------- History endpoints ----------
class TestHistory:
    def test_history_list(self, admin_headers):
        r = requests.get(f"{API}/gst/items/history", headers=admin_headers, timeout=15)
        assert r.status_code == 200
        body = r.json()
        assert "rows" in body and "count" in body
        assert isinstance(body["rows"], list)
        assert body["count"] >= 1

    def test_history_date_filter(self, admin_headers):
        r = requests.get(f"{API}/gst/items/history",
                         params={"start_date": "2026-01-01", "end_date": "2026-12-31"},
                         headers=admin_headers, timeout=15)
        assert r.status_code == 200
        rows = r.json().get("rows", [])
        for row in rows:
            assert "2026-01-01" <= row.get("effective_from", "") <= "2026-12-31"

    def test_history_excel_export(self, admin_headers):
        r = requests.get(f"{API}/gst/items/history/excel", headers=admin_headers, timeout=30)
        assert r.status_code == 200, r.text
        ctype = r.headers.get("content-type", "")
        assert "spreadsheetml" in ctype or "octet-stream" in ctype
        assert len(r.content) > 500
        wb = load_workbook(io.BytesIO(r.content), data_only=True)
        ws = wb.active
        # Header row is row 5 in the layout (rows 1-3 = company branding, row 4 = spacer)
        hdr = [str(c.value or "").strip() for c in ws[5]]
        assert "Item Name" in hdr and "Previous Rate" in hdr and "Updated Rate" in hdr

    def test_history_requires_admin(self):
        r = requests.get(f"{API}/gst/items/history", timeout=15)
        assert r.status_code in (401, 403)
        r = requests.get(f"{API}/gst/items/history/excel", timeout=15)
        assert r.status_code in (401, 403)


# ---------- Regression ----------
class TestRegression:
    def test_get_items(self, admin_headers):
        r = requests.get(f"{API}/gst/items", headers=admin_headers, timeout=15)
        assert r.status_code == 200

    def test_get_plans(self, admin_headers):
        r = requests.get(f"{API}/gst/plans", headers=admin_headers, timeout=15)
        assert r.status_code == 200

    def test_get_invoices(self, admin_headers):
        r = requests.get(f"{API}/gst/invoices",
                         params={"page": 1, "limit": 20}, headers=admin_headers, timeout=15)
        assert r.status_code == 200

    def test_reports_list(self, admin_headers):
        # Frontend calls /gst/reports/list
        r1 = requests.get(f"{API}/gst/reports/list", headers=admin_headers, timeout=15)
        # Spec mentions /gst-reports/list alternate prefix
        r2 = requests.get(f"{API}/gst-reports/list", headers=admin_headers, timeout=15)
        # At least one of these must respond 200
        assert (r1.status_code == 200) or (r2.status_code == 200), \
            f"both reports endpoints failed: {r1.status_code} / {r2.status_code}"
