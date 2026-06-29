"""
GST Billing Phase A - end-to-end backend tests.
Covers: branding/bank/T&C config fields, plant-warehouse exclusion,
Excel 32-col item-wise export with frozen panes + autofilter, PDF
Tax-Invoice format, amount-in-words helper.
"""
import os
import io
import time
import uuid
import pytest
import requests
from datetime import datetime
from openpyxl import load_workbook

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "").rstrip("/")
API = f"{BASE_URL}/api"

ADMIN = {"email": "admin@k3gas.com", "password": "Admin@123"}
HOLLONGI = {"email": "hollongi@k3gas.com", "password": "Hollongi@123"}


def _login(creds):
    r = requests.post(f"{API}/auth/login", json=creds, timeout=30)
    assert r.status_code == 200, f"login failed: {r.status_code} {r.text}"
    data = r.json()
    return data.get("token") or data.get("access_token")


@pytest.fixture(scope="module")
def admin_headers():
    tok = _login(ADMIN)
    return {"Authorization": f"Bearer {tok}", "Content-Type": "application/json"}


@pytest.fixture(scope="module")
def hollongi_headers():
    tok = _login(HOLLONGI)
    return {"Authorization": f"Bearer {tok}", "Content-Type": "application/json"}


# ---------- CONFIG NEW FIELDS ----------
class TestConfigNewFields:
    def test_get_config_has_all_new_fields(self, admin_headers):
        r = requests.get(f"{API}/gst/config", headers=admin_headers, timeout=30)
        assert r.status_code == 200
        cfg = r.json()
        expected_keys = [
            "company_name", "company_tagline", "company_address", "company_gstin",
            "company_state", "company_state_code", "company_phone", "company_email",
            "company_logo_url",
            "bank_name", "bank_account_no", "bank_ifsc", "bank_branch", "bank_account_holder",
            "terms_conditions", "signatory_name", "signatory_designation",
        ]
        for k in expected_keys:
            assert k in cfg, f"missing key {k} in gst_config"
        # terms_conditions should be present (default has 3 lines, but may be overwritten by other tests)
        assert "terms_conditions" in cfg
        assert isinstance(cfg.get("terms_conditions"), str)

    def test_put_config_persists_all_new_fields(self, admin_headers):
        payload = {
            "company_name": "TEST K3 GAS SERVICE",
            "company_tagline": "Reliable LPG Distribution",
            "company_address": "Test Address, Itanagar",
            "company_gstin": "12ABCDE1234F1Z5",
            "company_state": "Arunachal Pradesh",
            "company_state_code": "12",
            "company_phone": "9876543210",
            "company_email": "test@k3gas.com",
            "company_logo_url": "https://example.com/logo.png",
            "bank_name": "Test Bank Ltd",
            "bank_account_no": "1234567890",
            "bank_ifsc": "TEST0001234",
            "bank_branch": "Itanagar Branch",
            "bank_account_holder": "K3 Gas Service",
            "terms_conditions": "1. Test term one\n2. Test term two",
            "signatory_name": "Test Admin",
            "signatory_designation": "Director",
        }
        r = requests.put(f"{API}/gst/config", json=payload, headers=admin_headers, timeout=30)
        assert r.status_code == 200, r.text
        # Read back
        r2 = requests.get(f"{API}/gst/config", headers=admin_headers, timeout=30)
        cfg = r2.json()
        for k, v in payload.items():
            assert cfg.get(k) == v, f"{k}: expected {v!r}, got {cfg.get(k)!r}"


# ---------- PLANT WAREHOUSE EXCLUSION ----------
@pytest.fixture(scope="module")
def hollongi_warehouse_id(hollongi_headers):
    """Return warehouse id for Plant Hollongi (best effort via user info)."""
    r = requests.get(f"{API}/auth/me", headers=hollongi_headers, timeout=30)
    if r.status_code == 200:
        u = r.json()
        wid = u.get("warehouse_id") or u.get("warehouse", {}).get("id")
        if wid:
            return wid
    # Fallback: list warehouses (admin route)
    return None


class TestPlantWarehouseExclusion:
    def test_plant_warehouse_sale_skips_invoice(self, admin_headers, hollongi_headers):
        # Get hollongi user info
        me = requests.get(f"{API}/auth/me", headers=hollongi_headers, timeout=30).json()
        wh_id = me.get("warehouse_id")
        wh_name = me.get("warehouse_name") or me.get("warehouse", {}).get("name") or ""
        if not wh_id:
            # Try admin list
            wlist = requests.get(f"{API}/warehouses", headers=admin_headers, timeout=30)
            if wlist.status_code == 200:
                for w in wlist.json():
                    if "hollongi" in (w.get("name") or "").lower() or "plant" in (w.get("name") or "").lower():
                        wh_id = w.get("id")
                        wh_name = w.get("name")
                        break
        assert wh_id, "Could not resolve Plant Hollongi warehouse id"
        assert "hollongi" in wh_name.lower() or "plant" in wh_name.lower(), \
            f"Expected plant/hollongi in warehouse name, got {wh_name!r}"

        # Create a sales entry on plant warehouse via hollongi user
        cust = f"TEST_PlantCust_{uuid.uuid4().hex[:6]}"
        sale_payload = {
            "date": datetime.now().strftime("%Y-%m-%d"),
            "consumer_name": cust,
            "consumer_no": "9999900000",
            "address": "Plant test",
            "memo_no": f"TEST_{uuid.uuid4().hex[:6]}",
            "amount": 1100,
            "connection_type": "refill",
            "cylinder_nos": "1",
            "payment_mode": "cash",
            "cash_amount": 1100,
        }
        r = requests.post(f"{API}/sales-entries", json=sale_payload, headers=hollongi_headers, timeout=30)
        # Sale may succeed or fail based on schema - what matters is invoice NOT created
        if r.status_code not in (200, 201):
            pytest.skip(f"Plant sale-entry creation not permitted by API: {r.status_code} {r.text[:200]}")
        sale = r.json()
        sale_id = sale.get("id") or sale.get("_id") or sale.get("sale_id")
        assert sale_id, f"Could not extract sale id from response: {sale}"

        # Give async hook a moment
        time.sleep(1)

        # Query invoices for this sale via admin
        inv_list = requests.get(
            f"{API}/gst/invoices?search={cust}",
            headers=admin_headers, timeout=30,
        )
        assert inv_list.status_code == 200
        invs = inv_list.json()
        if isinstance(invs, dict):
            invs = invs.get("invoices") or invs.get("data") or []
        matched = [i for i in invs if i.get("sale_id") == sale_id]
        assert len(matched) == 0, \
            f"Plant warehouse sale should NOT auto-generate invoice; found: {matched}"


# ---------- EXCEL EXPORT 32-COLUMN ----------
class TestExcelExport:
    def test_excel_export_structure(self, admin_headers):
        r = requests.get(f"{API}/gst/invoices/export/excel", headers=admin_headers, timeout=60)
        assert r.status_code == 200, r.text[:300]
        assert "spreadsheet" in r.headers.get("content-type", "") or \
               "openxml" in r.headers.get("content-type", "")
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb.active

        # Title rows
        assert ws["A1"].value, "A1 (company name) should not be empty"
        a2 = (ws["A2"].value or "")
        assert "GSTIN" in a2, f"A2 should contain GSTIN, got {a2!r}"
        a3 = (ws["A3"].value or "")
        assert "Invoice Register" in a3, f"A3 should contain 'Invoice Register', got {a3!r}"
        assert "Generated" in a3

        # Header at row 5
        headers_row = [ws.cell(row=5, column=c).value for c in range(1, 33)]
        expected = [
            "Sl No.", "Invoice Number", "Invoice Date", "Status", "Customer Name", "Customer Mobile",
            "Customer GSTIN", "Item Name", "HSN Code", "Quantity", "Unit", "Rate",
            "Taxable Value", "GST %", "CGST %", "CGST Amount", "SGST %", "SGST Amount",
            "IGST %", "IGST Amount", "Total GST", "Discount", "Round Off", "Grand Total",
            "Payment Mode", "Warehouse", "Sales Executive", "Created By", "Created Date",
            "Cancelled By", "Cancelled Date", "Cancellation Reason",
        ]
        assert headers_row == expected, f"Header mismatch.\nGot: {headers_row}\nExpected: {expected}"

        # Frozen panes at A6
        assert ws.freeze_panes == "A6", f"freeze_panes should be A6, got {ws.freeze_panes!r}"

        # Autofilter
        assert ws.auto_filter.ref is not None, "Autofilter should be set"
        assert ws.auto_filter.ref.startswith("A5:AF"), \
            f"Autofilter should start at A5:AF, got {ws.auto_filter.ref}"


# ---------- PDF EXPORT ----------
class TestPdfExport:
    def _get_active_invoice_id(self, admin_headers):
        r = requests.get(f"{API}/gst/invoices?status=active", headers=admin_headers, timeout=30)
        assert r.status_code == 200
        data = r.json()
        invs = data if isinstance(data, list) else (data.get("invoices") or data.get("data") or [])
        if not invs:
            pytest.skip("No active invoices to test PDF export")
        return invs[0].get("id")

    def test_pdf_export_valid(self, admin_headers):
        inv_id = self._get_active_invoice_id(admin_headers)
        r = requests.get(f"{API}/gst/invoices/{inv_id}/pdf", headers=admin_headers, timeout=60)
        assert r.status_code == 200, r.text[:300]
        assert r.content[:4] == b"%PDF", f"Not a PDF, starts with {r.content[:8]!r}"
        assert len(r.content) > 1500, f"PDF too small ({len(r.content)} bytes), likely empty"


# ---------- AMOUNT IN WORDS ----------
class TestAmountInWords:
    """Direct import test of the helper."""
    def test_amount_in_words_samples(self):
        """Backend helper uses num2words(lang='en_IN') which produces text with
        'And' connectors and hyphens (e.g. 'Thirty-Four'). This differs from the
        frontend JS amountInWords helper (which omits 'And' and dashes). Both
        produce valid Indian-format words; we assert the backend variant here
        and report the divergence separately in the test report."""
        import sys
        sys.path.insert(0, "/app/backend")
        from routes.gst_billing import amount_in_words_inr
        r1 = amount_in_words_inr(100)
        assert r1.startswith("Rupees") and "One Hundred" in r1 and r1.endswith("Only"), f"Got {r1!r}"
        r2 = amount_in_words_inr(1234.56)
        # backend num2words en_IN -> "One Thousand Two Hundred And Thirty-Four"
        assert "Thirty" in r2 and "Four" in r2 and "Fifty" in r2 and "Six Paise" in r2 and r2.endswith("Only"), f"Got {r2!r}"
        r3 = amount_in_words_inr(123456)
        # backend produces 'Twenty-Three Thousand' with hyphen
        normalized = r3.replace("-", " ")
        assert "Lakh" in r3 and "Twenty Three Thousand" in normalized and r3.endswith("Only"), f"Got {r3!r}"
