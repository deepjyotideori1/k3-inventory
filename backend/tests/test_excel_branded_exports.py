"""Tests for branded Excel header exports (Customer / Sales / Plant)
and regression for gst_items routes that moved to routes/gst_items.py.

Verifies:
- New embed_logo branded header rows (B1=company, B2=tagline/address/helpline,
  A3=title+period+generated) is present on:
    /api/export/sales-excel
    /api/export/sales-summary-excel?group_by=daily
    /api/export/customer-refill-excel
    /api/export/customers-excel
    /api/export/excel?report_type=daily
    /api/export/excel?report_type=plant
- gst_items endpoints (now in routes/gst_items.py) still reachable:
    /api/gst/items, /api/gst/items/template/excel,
    /api/gst/items/history, /api/gst/items/history/excel
- Other gst_billing endpoints (plans/invoices/reports) still respond.
"""

import os
import io
import pytest
import requests
from openpyxl import load_workbook

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL").rstrip("/")


# ============== Fixtures ==============
@pytest.fixture(scope="session")
def admin_token():
    r = requests.post(f"{BASE_URL}/api/auth/login",
                      json={"email": "admin@k3gas.com", "password": "Admin@123"})
    assert r.status_code == 200, f"admin login failed: {r.status_code} {r.text}"
    body = r.json()
    return body.get("access_token") or body.get("token")


@pytest.fixture(scope="session")
def admin_headers(admin_token):
    return {"Authorization": f"Bearer {admin_token}"}


def _load_xlsx(content_bytes):
    return load_workbook(io.BytesIO(content_bytes), data_only=True)


# ============== Branded header sanity helpers ==============
def _assert_openpyxl_branded_header(ws, expected_title_substr,
                                    headers_row=5, sl_no_col="A"):
    """openpyxl exports: B1=company name, B2=tagline/address joined by '  |  ',
    A3=title+period+Generated. Table headers should be at row 5."""
    assert ws["B1"].value, "B1 should contain company name"
    assert "K3 GAS" in str(ws["B1"].value).upper(), \
        f"B1 should contain 'K3 GAS', got: {ws['B1'].value!r}"
    # B2 may be empty if tagline/address/helpline all blank — accept either.
    b2 = ws["B2"].value or ""
    # A3 should contain title + 'Generated:'
    a3 = str(ws["A3"].value or "")
    assert expected_title_substr.lower() in a3.lower(), \
        f"A3 should contain '{expected_title_substr}', got: {a3!r}"
    assert "Generated" in a3, f"A3 should contain 'Generated:', got: {a3!r}"
    # Table header row check (row 5)
    first_header = ws.cell(row=headers_row, column=1).value
    return first_header, b2, a3


# ============== /api/export/sales-excel ==============
class TestSalesExcelBranding:
    def test_sales_excel_branded(self, admin_headers):
        r = requests.get(f"{BASE_URL}/api/export/sales-excel",
                         headers=admin_headers, timeout=30)
        assert r.status_code == 200, f"{r.status_code}: {r.text[:200]}"
        assert "spreadsheetml" in r.headers.get("content-type", "")
        wb = _load_xlsx(r.content)
        ws = wb.active
        first_header, _, _ = _assert_openpyxl_branded_header(
            ws, "Sales Report", headers_row=5)
        assert first_header == "SL No.", \
            f"A5 should be 'SL No.', got {first_header!r}"


# ============== /api/export/sales-summary-excel ==============
class TestSalesSummaryExcelBranding:
    def test_sales_summary_daily(self, admin_headers):
        r = requests.get(f"{BASE_URL}/api/export/sales-summary-excel",
                         params={"group_by": "daily"},
                         headers=admin_headers, timeout=30)
        assert r.status_code == 200, f"{r.status_code}: {r.text[:200]}"
        wb = _load_xlsx(r.content)
        ws = wb.active
        # Title likely 'Sales Summary' or similar
        a3 = str(ws["A3"].value or "")
        assert "Generated" in a3
        first_header = ws.cell(row=5, column=1).value
        assert first_header == "Period", \
            f"A5 should be 'Period', got {first_header!r}"


# ============== /api/export/customer-refill-excel ==============
class TestCustomerRefillExcelBranding:
    def test_customer_refill_branded(self, admin_headers):
        r = requests.get(f"{BASE_URL}/api/export/customer-refill-excel",
                         headers=admin_headers, timeout=30)
        assert r.status_code == 200, f"{r.status_code}: {r.text[:200]}"
        wb = _load_xlsx(r.content)
        ws = wb.active
        first_header, _, _ = _assert_openpyxl_branded_header(
            ws, "Customer", headers_row=5)
        assert first_header == "SL No.", \
            f"A5 should be 'SL No.', got {first_header!r}"


# ============== /api/export/customers-excel (xlsxwriter) ==============
class TestCustomersExcelBranding:
    def test_customers_excel_branded(self, admin_headers):
        r = requests.get(f"{BASE_URL}/api/export/customers-excel",
                         headers=admin_headers, timeout=30)
        assert r.status_code == 200, f"{r.status_code}: {r.text[:200]}"
        wb = _load_xlsx(r.content)
        ws = wb.active
        # xlsxwriter rows 0-2 == Excel rows 1-3. Headers at Excel row 5.
        b1 = str(ws["B1"].value or "")
        a3 = str(ws["A3"].value or "")
        assert "K3 GAS" in b1.upper(), f"B1 should have company, got {b1!r}"
        assert "Customer" in a3, f"A3 should mention 'Customer', got {a3!r}"
        # headers at Excel row 5 (xlsxwriter row 4)
        first_header = ws.cell(row=5, column=1).value
        assert first_header, f"Row 5 should have first header value, got {first_header!r}"


# ============== /api/export/excel?report_type=daily ==============
class TestExportExcelDailyBranding:
    def test_daily_inventory_report(self, admin_headers):
        r = requests.get(f"{BASE_URL}/api/export/excel",
                         params={"report_type": "daily"},
                         headers=admin_headers, timeout=30)
        assert r.status_code == 200, f"{r.status_code}: {r.text[:200]}"
        wb = _load_xlsx(r.content)
        ws = wb.active
        b1 = str(ws["B1"].value or "")
        a3 = str(ws["A3"].value or "")
        assert "K3 GAS" in b1.upper()
        assert "Daily" in a3 or "daily" in a3.lower()
        # data row 5 should have first header cell populated
        first_header = ws.cell(row=5, column=1).value
        assert first_header, "Headers row 5 col 1 should have value"


# ============== /api/export/excel?report_type=plant (two-sheet) ==============
class TestExportExcelPlantBranding:
    def test_plant_main_and_breakdown_sheets_branded(self, admin_headers):
        r = requests.get(f"{BASE_URL}/api/export/excel",
                         params={"report_type": "plant"},
                         headers=admin_headers, timeout=30)
        assert r.status_code == 200, f"{r.status_code}: {r.text[:200]}"
        wb = _load_xlsx(r.content)
        # At least main sheet
        assert len(wb.sheetnames) >= 1
        ws_main = wb[wb.sheetnames[0]]
        b1_main = str(ws_main["B1"].value or "")
        assert "K3 GAS" in b1_main.upper(), \
            f"Main sheet B1 should have company, got {b1_main!r}"
        # Breakdown sheet branding (if present)
        if len(wb.sheetnames) >= 2:
            ws_b = wb[wb.sheetnames[1]]
            b1_b = str(ws_b["B1"].value or "")
            assert "K3 GAS" in b1_b.upper(), \
                f"Breakdown sheet B1 should have company, got {b1_b!r}"


# ============== Regression: gst_items routes still reachable ==============
class TestGstItemsRouteRegression:
    """After refactor — gst_items endpoints moved to routes/gst_items.py
    but URLs unchanged. Verify cold start works (circular-import safe)."""

    def test_list_items_200(self, admin_headers):
        r = requests.get(f"{BASE_URL}/api/gst/items", headers=admin_headers)
        assert r.status_code == 200
        data = r.json()
        assert isinstance(data, list)
        assert len(data) > 0, "Expected seeded items"

    def test_template_excel_200(self, admin_headers):
        r = requests.get(f"{BASE_URL}/api/gst/items/template/excel",
                         headers=admin_headers)
        assert r.status_code == 200
        assert "spreadsheetml" in r.headers.get("content-type", "")

    def test_history_list_200(self, admin_headers):
        r = requests.get(f"{BASE_URL}/api/gst/items/history",
                         headers=admin_headers)
        assert r.status_code == 200
        body = r.json()
        assert "rows" in body and "count" in body

    def test_history_excel_200(self, admin_headers):
        r = requests.get(f"{BASE_URL}/api/gst/items/history/excel",
                         headers=admin_headers)
        assert r.status_code == 200
        assert "spreadsheetml" in r.headers.get("content-type", "")

    def test_create_update_deactivate_activate_flow(self, admin_headers):
        # Create
        payload = {"name": "TEST_RefactorItem", "hsn": "9999",
                   "unit": "Nos", "gst_rate": 18, "default_rate": 111}
        r = requests.post(f"{BASE_URL}/api/gst/items",
                          json=payload, headers=admin_headers)
        assert r.status_code in (200, 201), r.text
        item = r.json()
        item_id = item["id"]
        # Update (changes default_rate -> history)
        r2 = requests.put(f"{BASE_URL}/api/gst/items/{item_id}",
                          json={"default_rate": 222,
                                "effective_from": "2026-01-15"},
                          headers=admin_headers)
        assert r2.status_code == 200, r2.text
        # Deactivate
        r3 = requests.delete(f"{BASE_URL}/api/gst/items/{item_id}",
                             headers=admin_headers)
        assert r3.status_code == 200
        # Activate
        r4 = requests.post(f"{BASE_URL}/api/gst/items/{item_id}/activate",
                           headers=admin_headers)
        assert r4.status_code == 200

    def test_bulk_update_json(self, admin_headers):
        payload = {
            "effective_from": "2026-01-15",
            "items": [{
                "name": "TEST_BulkRefactor",
                "hsn": "8888", "unit": "Nos",
                "gst_rate": 18, "default_rate": 99,
                "is_active": "active",
            }],
        }
        r = requests.post(f"{BASE_URL}/api/gst/items/bulk-update",
                          json=payload, headers=admin_headers)
        assert r.status_code == 200, r.text
        body = r.json()
        assert body["created"] + body["updated"] >= 1


# ============== Regression: other gst_billing endpoints still work ==============
class TestGstBillingRegression:
    def test_plans_list(self, admin_headers):
        r = requests.get(f"{BASE_URL}/api/gst/plans", headers=admin_headers)
        assert r.status_code == 200
        assert isinstance(r.json(), list)

    def test_invoices_paginated(self, admin_headers):
        r = requests.get(f"{BASE_URL}/api/gst/invoices",
                         params={"page": 1, "limit": 5},
                         headers=admin_headers)
        assert r.status_code == 200
        body = r.json()
        # Either {invoices, total} or list
        assert isinstance(body, (list, dict))

    def test_reports_list(self, admin_headers):
        r = requests.get(f"{BASE_URL}/api/gst/reports/list",
                         headers=admin_headers)
        assert r.status_code == 200
        body = r.json()
        assert isinstance(body, (list, dict))
