"""
Test HRMS Employee Bulk Upload Feature
Tests: Template download, validation, confirm import (skip/overwrite), error report, upload logs
"""
import pytest
import requests
import os
from io import BytesIO

BASE_URL = os.environ.get('REACT_APP_BACKEND_URL', '').rstrip('/')

# Test credentials
ADMIN_EMAIL = "admin@k3gas.com"
ADMIN_PASSWORD = "Admin@123"


@pytest.fixture(scope="module")
def auth_token():
    """Get admin auth token"""
    response = requests.post(f"{BASE_URL}/api/auth/login", json={
        "email": ADMIN_EMAIL,
        "password": ADMIN_PASSWORD
    })
    if response.status_code == 200:
        return response.json().get("token")
    pytest.skip(f"Authentication failed: {response.status_code} - {response.text}")


@pytest.fixture(scope="module")
def auth_headers(auth_token):
    """Auth headers for requests"""
    return {"Authorization": f"Bearer {auth_token}"}


@pytest.fixture(scope="module")
def departments(auth_headers):
    """Get existing departments"""
    response = requests.get(f"{BASE_URL}/api/hrms/departments", headers=auth_headers)
    if response.status_code == 200:
        return response.json()
    return []


def create_test_xlsx(rows_data, include_headers=True):
    """Create a test xlsx file with given data using openpyxl"""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    
    headers = [
        'Employee_ID', 'Full_Name', 'Gender', 'Date_of_Birth', 'Phone', 'Email',
        'Address', 'Department', 'Designation', 'Date_of_Joining', 'Employment_Type',
        'Basic_Salary', 'HRA', 'Allowances', 'Bank_Account_Number', 'IFSC_Code',
        'PAN_Number', 'Aadhaar_Number', 'PF_Applicable', 'ESI_Applicable'
    ]
    
    if include_headers:
        for col, h in enumerate(headers, 1):
            ws.cell(row=1, column=col, value=h)
        start_row = 2
    else:
        start_row = 1
    
    for row_idx, row_data in enumerate(rows_data, start_row):
        for col_idx, val in enumerate(row_data, 1):
            ws.cell(row=row_idx, column=col_idx, value=val)
    
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


class TestBulkUploadTemplate:
    """Test template download endpoint"""
    
    def test_download_template_success(self, auth_headers):
        """GET /api/hrms/employees/bulk-upload/template - downloads valid xlsx"""
        response = requests.get(
            f"{BASE_URL}/api/hrms/employees/bulk-upload/template",
            headers=auth_headers
        )
        assert response.status_code == 200, f"Expected 200, got {response.status_code}"
        assert 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet' in response.headers.get('Content-Type', '')
        assert 'attachment' in response.headers.get('Content-Disposition', '')
        assert len(response.content) > 1000, "Template file too small"
        print("✓ Template download successful - valid xlsx file returned")
    
    def test_download_template_without_auth(self):
        """Template download requires authentication"""
        response = requests.get(f"{BASE_URL}/api/hrms/employees/bulk-upload/template")
        assert response.status_code in [401, 403], f"Expected 401/403, got {response.status_code}"
        print("✓ Template download correctly requires authentication")


class TestBulkUploadValidation:
    """Test validation endpoint"""
    
    def test_validate_valid_file(self, auth_headers, departments):
        """POST /api/hrms/employees/bulk-upload/validate - validates valid xlsx"""
        dept_name = departments[0]['name'] if departments else 'Operations'
        
        rows = [
            ['TEST-BU-001', 'John Test Smith', 'Male', '15-06-1990', '9876543210', 'john.test@example.com',
             'Test Address', dept_name, 'Manager', '01-01-2024', 'Full-time',
             25000, 5000, 3000, '1234567890', 'SBIN0001234',
             'ABCDE1234F', '123456789012', 'Yes', 'Yes'],
        ]
        xlsx_file = create_test_xlsx(rows)
        
        response = requests.post(
            f"{BASE_URL}/api/hrms/employees/bulk-upload/validate",
            headers=auth_headers,
            files={'file': ('test_upload.xlsx', xlsx_file, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}
        )
        assert response.status_code == 200, f"Expected 200, got {response.status_code}: {response.text}"
        data = response.json()
        
        assert 'total_rows' in data
        assert 'valid_count' in data
        assert 'error_count' in data
        assert 'new_count' in data
        assert 'update_count' in data
        assert 'rows' in data
        assert data['total_rows'] >= 1
        print(f"✓ Validation successful - {data['total_rows']} rows, {data['valid_count']} valid, {data['error_count']} errors")
    
    def test_validate_missing_required_fields(self, auth_headers):
        """Validation detects missing required fields"""
        rows = [
            ['', '', '', '', '', '', '', '', '', '', '', '', '', '', '', '', '', '', '', ''],  # All empty
            ['TEST-BU-002', '', '', '', '', '', '', '', '', '', '', '', '', '', '', '', '', '', '', ''],  # Missing name
            ['', 'Test Name', '', '', '', '', '', '', '', '01-01-2024', '', 25000, '', '', '', '', '', '', '', ''],  # Missing ID
        ]
        xlsx_file = create_test_xlsx(rows)
        
        response = requests.post(
            f"{BASE_URL}/api/hrms/employees/bulk-upload/validate",
            headers=auth_headers,
            files={'file': ('test_upload.xlsx', xlsx_file, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}
        )
        assert response.status_code == 200
        data = response.json()
        
        # Should have errors for missing required fields
        assert data['error_count'] >= 2, f"Expected at least 2 errors, got {data['error_count']}"
        
        # Check error messages
        errors_found = []
        for row in data['rows']:
            if row.get('has_errors'):
                errors_found.extend(row.get('errors', []))
        
        assert any('Employee_ID' in e or 'Full_Name' in e or 'required' in e.lower() for e in errors_found), \
            f"Expected required field errors, got: {errors_found}"
        print(f"✓ Missing required fields detected - {data['error_count']} errors found")
    
    def test_validate_invalid_pan_format(self, auth_headers, departments):
        """Validation detects invalid PAN format"""
        dept_name = departments[0]['name'] if departments else 'Operations'
        
        rows = [
            ['TEST-BU-003', 'Test Pan User', 'Male', '15-06-1990', '9876543210', 'pan.test@example.com',
             'Test Address', dept_name, 'Manager', '01-01-2024', 'Full-time',
             25000, 5000, 3000, '1234567890', 'SBIN0001234',
             'INVALID123', '123456789012', 'Yes', 'Yes'],  # Invalid PAN
        ]
        xlsx_file = create_test_xlsx(rows)
        
        response = requests.post(
            f"{BASE_URL}/api/hrms/employees/bulk-upload/validate",
            headers=auth_headers,
            files={'file': ('test_upload.xlsx', xlsx_file, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}
        )
        assert response.status_code == 200
        data = response.json()
        
        # Find the row with PAN error
        pan_error_found = False
        for row in data['rows']:
            if row.get('employee_id') == 'TEST-BU-003' and row.get('has_errors'):
                for err in row.get('errors', []):
                    if 'PAN' in err.upper():
                        pan_error_found = True
                        break
        
        assert pan_error_found, f"Expected PAN format error, got rows: {data['rows']}"
        print("✓ Invalid PAN format detected correctly")
    
    def test_validate_invalid_aadhaar_format(self, auth_headers, departments):
        """Validation detects invalid Aadhaar format"""
        dept_name = departments[0]['name'] if departments else 'Operations'
        
        rows = [
            ['TEST-BU-004', 'Test Aadhaar User', 'Male', '15-06-1990', '9876543210', 'aadhaar.test@example.com',
             'Test Address', dept_name, 'Manager', '01-01-2024', 'Full-time',
             25000, 5000, 3000, '1234567890', 'SBIN0001234',
             'ABCDE1234F', '12345', 'Yes', 'Yes'],  # Invalid Aadhaar (not 12 digits)
        ]
        xlsx_file = create_test_xlsx(rows)
        
        response = requests.post(
            f"{BASE_URL}/api/hrms/employees/bulk-upload/validate",
            headers=auth_headers,
            files={'file': ('test_upload.xlsx', xlsx_file, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}
        )
        assert response.status_code == 200
        data = response.json()
        
        # Find the row with Aadhaar error
        aadhaar_error_found = False
        for row in data['rows']:
            if row.get('employee_id') == 'TEST-BU-004' and row.get('has_errors'):
                for err in row.get('errors', []):
                    if 'Aadhaar' in err or 'aadhaar' in err.lower():
                        aadhaar_error_found = True
                        break
        
        assert aadhaar_error_found, f"Expected Aadhaar format error, got rows: {data['rows']}"
        print("✓ Invalid Aadhaar format detected correctly")
    
    def test_validate_invalid_date_format(self, auth_headers, departments):
        """Validation detects invalid date format"""
        dept_name = departments[0]['name'] if departments else 'Operations'
        
        rows = [
            ['TEST-BU-005', 'Test Date User', 'Male', '1990-06-15', '9876543210', 'date.test@example.com',
             'Test Address', dept_name, 'Manager', 'invalid-date', 'Full-time',
             25000, 5000, 3000, '1234567890', 'SBIN0001234',
             'ABCDE1234F', '123456789012', 'Yes', 'Yes'],  # Invalid DOJ
        ]
        xlsx_file = create_test_xlsx(rows)
        
        response = requests.post(
            f"{BASE_URL}/api/hrms/employees/bulk-upload/validate",
            headers=auth_headers,
            files={'file': ('test_upload.xlsx', xlsx_file, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}
        )
        assert response.status_code == 200
        data = response.json()
        
        # Find the row with date error
        date_error_found = False
        for row in data['rows']:
            if row.get('employee_id') == 'TEST-BU-005' and row.get('has_errors'):
                for err in row.get('errors', []):
                    if 'Date' in err or 'date' in err.lower():
                        date_error_found = True
                        break
        
        assert date_error_found, f"Expected date format error, got rows: {data['rows']}"
        print("✓ Invalid date format detected correctly")
    
    def test_validate_duplicate_ids_in_file(self, auth_headers, departments):
        """Validation detects duplicate Employee_IDs within the file"""
        dept_name = departments[0]['name'] if departments else 'Operations'
        
        rows = [
            ['TEST-BU-DUP', 'First User', 'Male', '15-06-1990', '9876543210', 'first@example.com',
             'Test Address', dept_name, 'Manager', '01-01-2024', 'Full-time',
             25000, 5000, 3000, '1234567890', 'SBIN0001234',
             'ABCDE1234F', '123456789012', 'Yes', 'Yes'],
            ['TEST-BU-DUP', 'Second User', 'Female', '20-07-1992', '9876543211', 'second@example.com',
             'Test Address 2', dept_name, 'Executive', '01-02-2024', 'Full-time',
             20000, 4000, 2000, '1234567891', 'SBIN0001235',
             'FGHIJ5678K', '234567890123', 'Yes', 'No'],  # Same ID
        ]
        xlsx_file = create_test_xlsx(rows)
        
        response = requests.post(
            f"{BASE_URL}/api/hrms/employees/bulk-upload/validate",
            headers=auth_headers,
            files={'file': ('test_upload.xlsx', xlsx_file, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}
        )
        assert response.status_code == 200
        data = response.json()
        
        # Should detect duplicate
        dup_error_found = False
        for row in data['rows']:
            if row.get('has_errors'):
                for err in row.get('errors', []):
                    if 'Duplicate' in err or 'duplicate' in err.lower():
                        dup_error_found = True
                        break
        
        assert dup_error_found, f"Expected duplicate ID error, got rows: {data['rows']}"
        print("✓ Duplicate Employee_ID in file detected correctly")
    
    def test_validate_auto_capitalize_names(self, auth_headers, departments):
        """Validation auto-capitalizes names"""
        dept_name = departments[0]['name'] if departments else 'Operations'
        
        rows = [
            ['TEST-BU-CAP', 'john doe smith', 'Male', '15-06-1990', '9876543210', 'cap.test@example.com',
             'Test Address', dept_name, 'Manager', '01-01-2024', 'Full-time',
             25000, 5000, 3000, '1234567890', 'SBIN0001234',
             'ABCDE1234F', '123456789012', 'Yes', 'Yes'],
        ]
        xlsx_file = create_test_xlsx(rows)
        
        response = requests.post(
            f"{BASE_URL}/api/hrms/employees/bulk-upload/validate",
            headers=auth_headers,
            files={'file': ('test_upload.xlsx', xlsx_file, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}
        )
        assert response.status_code == 200
        data = response.json()
        
        # Find the row and check name is capitalized
        for row in data['rows']:
            if row.get('employee_id') == 'TEST-BU-CAP':
                assert row.get('name') == 'John Doe Smith', f"Expected 'John Doe Smith', got '{row.get('name')}'"
                print("✓ Name auto-capitalization working correctly")
                return
        
        pytest.fail("Could not find test row for capitalization check")
    
    def test_validate_non_xlsx_file_rejected(self, auth_headers):
        """Validation rejects non-xlsx files"""
        # Create a fake CSV content
        csv_content = b"Employee_ID,Full_Name\nTEST-001,Test User"
        
        response = requests.post(
            f"{BASE_URL}/api/hrms/employees/bulk-upload/validate",
            headers=auth_headers,
            files={'file': ('test_upload.csv', csv_content, 'text/csv')}
        )
        assert response.status_code == 400, f"Expected 400, got {response.status_code}"
        print("✓ Non-xlsx file correctly rejected")
    
    def test_validate_without_auth(self):
        """Validation requires authentication"""
        rows = [['TEST-001', 'Test User', '', '', '', '', '', '', '', '01-01-2024', '', 25000, '', '', '', '', '', '', '', '']]
        xlsx_file = create_test_xlsx(rows)
        
        response = requests.post(
            f"{BASE_URL}/api/hrms/employees/bulk-upload/validate",
            files={'file': ('test_upload.xlsx', xlsx_file, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}
        )
        assert response.status_code in [401, 403], f"Expected 401/403, got {response.status_code}"
        print("✓ Validation correctly requires authentication")


class TestBulkUploadConfirm:
    """Test confirm import endpoint"""
    
    def test_confirm_skip_mode(self, auth_headers, departments):
        """POST /api/hrms/employees/bulk-upload/confirm with mode=skip"""
        dept_name = departments[0]['name'] if departments else 'Operations'
        dept_id = departments[0]['id'] if departments else ''
        
        # First validate
        rows = [
            ['TEST-BU-SKIP-001', 'Skip Test User', 'Male', '15-06-1990', '9876543210', 'skip.test@example.com',
             'Test Address', dept_name, 'Manager', '01-01-2024', 'Full-time',
             25000, 5000, 3000, '1234567890', 'SBIN0001234',
             'ABCDE1234F', '123456789012', 'Yes', 'Yes'],
        ]
        xlsx_file = create_test_xlsx(rows)
        
        validate_response = requests.post(
            f"{BASE_URL}/api/hrms/employees/bulk-upload/validate",
            headers=auth_headers,
            files={'file': ('test_upload.xlsx', xlsx_file, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}
        )
        assert validate_response.status_code == 200
        validate_data = validate_response.json()
        
        # Confirm with skip mode
        confirm_response = requests.post(
            f"{BASE_URL}/api/hrms/employees/bulk-upload/confirm",
            headers=auth_headers,
            json={
                'rows': validate_data['rows'],
                'mode': 'skip'
            }
        )
        assert confirm_response.status_code == 200, f"Expected 200, got {confirm_response.status_code}: {confirm_response.text}"
        confirm_data = confirm_response.json()
        
        assert 'created' in confirm_data
        assert 'updated' in confirm_data
        assert 'skipped' in confirm_data
        assert 'message' in confirm_data
        print(f"✓ Confirm (skip mode) successful - created: {confirm_data['created']}, updated: {confirm_data['updated']}, skipped: {confirm_data['skipped']}")
    
    def test_confirm_overwrite_mode(self, auth_headers, departments):
        """POST /api/hrms/employees/bulk-upload/confirm with mode=overwrite"""
        dept_name = departments[0]['name'] if departments else 'Operations'
        
        # Use same ID as skip test to test overwrite
        rows = [
            ['TEST-BU-SKIP-001', 'Updated Skip Test User', 'Male', '15-06-1990', '9876543211', 'skip.updated@example.com',
             'Updated Address', dept_name, 'Senior Manager', '01-01-2024', 'Full-time',
             30000, 6000, 4000, '1234567890', 'SBIN0001234',
             'ABCDE1234F', '123456789012', 'Yes', 'Yes'],
        ]
        xlsx_file = create_test_xlsx(rows)
        
        validate_response = requests.post(
            f"{BASE_URL}/api/hrms/employees/bulk-upload/validate",
            headers=auth_headers,
            files={'file': ('test_upload.xlsx', xlsx_file, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}
        )
        assert validate_response.status_code == 200
        validate_data = validate_response.json()
        
        # Should detect as existing
        assert validate_data['update_count'] >= 1, f"Expected at least 1 existing, got {validate_data['update_count']}"
        
        # Confirm with overwrite mode
        confirm_response = requests.post(
            f"{BASE_URL}/api/hrms/employees/bulk-upload/confirm",
            headers=auth_headers,
            json={
                'rows': validate_data['rows'],
                'mode': 'overwrite'
            }
        )
        assert confirm_response.status_code == 200
        confirm_data = confirm_response.json()
        
        assert confirm_data['updated'] >= 1, f"Expected at least 1 updated, got {confirm_data['updated']}"
        print(f"✓ Confirm (overwrite mode) successful - updated: {confirm_data['updated']}")
    
    def test_confirm_empty_rows_rejected(self, auth_headers):
        """Confirm with empty rows returns error"""
        response = requests.post(
            f"{BASE_URL}/api/hrms/employees/bulk-upload/confirm",
            headers=auth_headers,
            json={'rows': [], 'mode': 'skip'}
        )
        assert response.status_code == 400, f"Expected 400, got {response.status_code}"
        print("✓ Empty rows correctly rejected")
    
    def test_confirm_without_auth(self):
        """Confirm requires authentication"""
        response = requests.post(
            f"{BASE_URL}/api/hrms/employees/bulk-upload/confirm",
            json={'rows': [{'employee_id': 'TEST-001'}], 'mode': 'skip'}
        )
        assert response.status_code in [401, 403], f"Expected 401/403, got {response.status_code}"
        print("✓ Confirm correctly requires authentication")


class TestBulkUploadErrorReport:
    """Test error report download endpoint"""
    
    def test_download_error_report(self, auth_headers):
        """POST /api/hrms/employees/bulk-upload/error-report - generates xlsx"""
        errors = [
            {'row': 2, 'employee_id': 'TEST-ERR-001', 'name': 'Error User 1', 'errors': ['Invalid PAN format', 'Missing Date_of_Joining']},
            {'row': 3, 'employee_id': 'TEST-ERR-002', 'name': 'Error User 2', 'errors': ['Invalid Aadhaar']},
        ]
        
        response = requests.post(
            f"{BASE_URL}/api/hrms/employees/bulk-upload/error-report",
            headers=auth_headers,
            json={'errors': errors}
        )
        assert response.status_code == 200, f"Expected 200, got {response.status_code}"
        assert 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet' in response.headers.get('Content-Type', '')
        assert len(response.content) > 500, "Error report file too small"
        print("✓ Error report download successful - valid xlsx returned")
    
    def test_error_report_without_auth(self):
        """Error report requires authentication"""
        response = requests.post(
            f"{BASE_URL}/api/hrms/employees/bulk-upload/error-report",
            json={'errors': []}
        )
        assert response.status_code in [401, 403], f"Expected 401/403, got {response.status_code}"
        print("✓ Error report correctly requires authentication")


class TestUploadLogs:
    """Test upload logs endpoint"""
    
    def test_get_upload_logs(self, auth_headers):
        """GET /api/hrms/employees/upload-logs - returns upload history"""
        response = requests.get(
            f"{BASE_URL}/api/hrms/employees/upload-logs",
            headers=auth_headers
        )
        assert response.status_code == 200, f"Expected 200, got {response.status_code}"
        data = response.json()
        
        assert isinstance(data, list), "Expected list of logs"
        
        # If there are logs, verify structure
        if len(data) > 0:
            log = data[0]
            assert 'id' in log
            assert 'uploaded_by' in log
            assert 'mode' in log
            assert 'total_rows' in log
            assert 'created' in log
            assert 'updated' in log
            assert 'skipped' in log
            assert 'created_at' in log
        
        print(f"✓ Upload logs retrieved - {len(data)} logs found")
    
    def test_upload_logs_without_auth(self):
        """Upload logs requires authentication"""
        response = requests.get(f"{BASE_URL}/api/hrms/employees/upload-logs")
        assert response.status_code in [401, 403], f"Expected 401/403, got {response.status_code}"
        print("✓ Upload logs correctly requires authentication")


class TestEmployeeCRUDRegression:
    """Regression tests for existing employee CRUD"""
    
    def test_get_employees(self, auth_headers):
        """GET /api/hrms/employees - still works"""
        response = requests.get(
            f"{BASE_URL}/api/hrms/employees",
            headers=auth_headers
        )
        assert response.status_code == 200, f"Expected 200, got {response.status_code}"
        data = response.json()
        assert 'employees' in data
        assert 'total' in data
        print(f"✓ GET employees working - {data['total']} employees found")
    
    def test_get_departments(self, auth_headers):
        """GET /api/hrms/departments - still works"""
        response = requests.get(
            f"{BASE_URL}/api/hrms/departments",
            headers=auth_headers
        )
        assert response.status_code == 200, f"Expected 200, got {response.status_code}"
        data = response.json()
        assert isinstance(data, list)
        print(f"✓ GET departments working - {len(data)} departments found")


class TestCleanup:
    """Cleanup test data"""
    
    def test_cleanup_test_employees(self, auth_headers):
        """Remove test employees created during tests"""
        # Get all employees
        response = requests.get(
            f"{BASE_URL}/api/hrms/employees?limit=100",
            headers=auth_headers
        )
        if response.status_code == 200:
            employees = response.json().get('employees', [])
            deleted = 0
            for emp in employees:
                if emp.get('employee_id', '').startswith('TEST-BU-'):
                    # Delete test employee
                    del_response = requests.delete(
                        f"{BASE_URL}/api/hrms/employees/{emp['id']}",
                        headers=auth_headers
                    )
                    if del_response.status_code in [200, 204]:
                        deleted += 1
            print(f"✓ Cleanup complete - {deleted} test employees deactivated")
        else:
            print("⚠ Could not fetch employees for cleanup")


if __name__ == "__main__":
    pytest.main([__file__, "-v", "--tb=short"])
